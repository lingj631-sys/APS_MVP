"""ACT-25/26: Excel validation, idempotent import, versioning, rollback."""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from src.rule_sets import excel as rule_excel
from src.rule_sets import service


class TestRuleSetPlatform(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="aps_rules_test_"))
        cls.db = cls.tmp / "rules.sqlite3"
        cls.created = service.ensure_seed(("WL",), db_path=cls.db)

    def test_01_seed_and_active(self):
        self.assertEqual(self.created, ["WL"])
        rule_set, warnings = service.active_rule_set("WL", db_path=self.db)
        self.assertIsNotNone(rule_set)
        self.assertEqual(rule_set.rule_set_id, "WL_PHASE1")
        self.assertEqual(len(rule_set.rules), 13)
        ids = {rule.rule_id for rule in rule_set.rules}
        self.assertTrue({"WL_MPQ_FG", "WL_MPQ_SF", "WL_CAP_FG", "WL_CAP_SF"} <= ids)
        self.assertEqual(len(rule_set.groups), 5)
        self.assertIsInstance(warnings, list)

    def test_02_template_roundtrip_is_idempotent(self):
        template = self.tmp / "template.xlsx"
        rule_excel.write_template(template, site="WL", rule_set_id="WL_ROUNDTRIP")
        first = service.import_excel(template, site="WL", db_path=self.db)
        self.assertTrue(first["ok"], first)
        self.assertTrue(first["created"])
        second = service.import_excel(template, site="WL", db_path=self.db)
        self.assertTrue(second["ok"], second)
        self.assertFalse(second["created"])
        self.assertEqual(second["version"], first["version"])

    def test_03_bad_excel_reports_cell(self):
        template = self.tmp / "bad.xlsx"
        rule_excel.write_template(template, site="WL", rule_set_id="WL_BAD")
        workbook = load_workbook(template)
        sheet = workbook["Rules"]
        sheet["G2"] = "sometimes"
        workbook.save(template)
        payload, errors, _ = rule_excel.read_workbook(template)
        self.assertIsNone(payload)
        self.assertTrue(any(error["cell"].endswith("G2") for error in errors), errors)

    def test_04_bad_template_reports_cell(self):
        template = self.tmp / "bad_template.xlsx"
        rule_excel.write_template(template, site="WL", rule_set_id="WL_BAD2")
        workbook = load_workbook(template)
        sheet = workbook["Rules"]
        sheet["B2"] = "no_such_template"
        workbook.save(template)
        payload, errors, _ = rule_excel.read_workbook(template)
        self.assertIsNone(payload)
        self.assertTrue(any(error["cell"].endswith("B2") and "未知模板" in error["message"] for error in errors), errors)

    def test_05_version_diff_and_rollback(self):
        rule_set, _ = service.active_rule_set("WL", db_path=self.db)
        payload = rule_set.to_dict()
        payload["rules"][0]["weight"] = 321
        saved = service.save_payload(payload, site="WL", note="weight", activate=True, db_path=self.db)
        self.assertTrue(saved["ok"], saved)
        self.assertTrue(saved["created"])
        self.assertEqual(saved["diff"]["summary"]["rules_modified"], 1)
        active, _ = service.active_rule_set("WL", db_path=self.db)
        self.assertEqual(active.rules[0].weight, 321)

        first_version = [v for v in service.list_versions("WL", db_path=self.db) if v["status"] == "archived"][-1]
        service.activate(first_version["id"], db_path=self.db)
        restored, _ = service.active_rule_set("WL", db_path=self.db)
        self.assertNotEqual(restored.rules[0].weight, 321)

    def test_06_save_rejects_invalid_payload(self):
        rule_set, _ = service.active_rule_set("WL", db_path=self.db)
        payload = rule_set.to_dict()
        payload["rules"][0]["template"] = "no_such_template"
        result = service.save_payload(payload, site="WL", db_path=self.db)
        self.assertFalse(result["ok"])
        self.assertTrue(result["errors"])

    @staticmethod
    def _toy_payload():
        return {
            "rule_set_id": "TOY",
            "site": "WL",
            "version": 1,
            "groups": [
                {"id": "G0", "priority": 0, "mode": "hard"},
                {"id": "G1", "priority": 10, "mode": "weighted"},
            ],
            "rules": [
                {"id": "R1", "template": "changeover", "scope": {}, "nature": "soft", "weight": 100, "group": "G1", "priority": 10, "params": {}},
                {"id": "R2", "template": "capacity_utilization", "scope": {}, "nature": "soft", "weight": 100, "group": "G1", "priority": 10, "params": {}},
            ],
        }

    def test_07_sensitivity_scales_soft_weights(self):
        from src.rule_sets import sensitivity

        payload = self._toy_payload()

        def solve(candidate):
            weight = float(candidate["rules"][0]["weight"])
            return {"feasible": True, "objective": weight, "blocks": 100 - weight / 10}

        report = sensitivity.run_sensitivity(payload, solve, pct=0.2)
        self.assertTrue(report["feasible"])
        self.assertEqual(len(report["rows"]), 2)
        by_direction = {row["direction"]: row for row in report["rows"]}
        self.assertAlmostEqual(by_direction["+20%"]["kpis"]["objective"], 120.0)
        self.assertAlmostEqual(by_direction["-20%"]["kpis"]["objective"], 80.0)
        self.assertAlmostEqual(by_direction["+20%"]["delta"]["objective"], 20.0)

    def test_08_diagnose_single_rule(self):
        from src.rule_sets import diagnose as diagnose_mod

        payload = self._toy_payload()

        def solve(candidate):
            return {"feasible": not any(rule["id"] == "R2" for rule in candidate["rules"])}

        report = diagnose_mod.diagnose(payload, solve)
        self.assertFalse(report["feasible"])
        self.assertIn("G1", report["conflict_groups"])
        self.assertIn("R2", report["conflict_rules"])
        self.assertIn("R2", report["message"])

    def test_09_diagnose_rule_pair(self):
        from src.rule_sets import diagnose as diagnose_mod

        payload = self._toy_payload()

        def solve(candidate):
            ids = {rule["id"] for rule in candidate["rules"]}
            return {"feasible": not ({"R1", "R2"} & ids)}

        report = diagnose_mod.diagnose(payload, solve)
        self.assertFalse(report["feasible"])
        self.assertIn("G1", report["conflict_groups"])
        self.assertEqual(report["conflict_rules"], [])
        self.assertEqual(report["conflict_pairs"], [["R1", "R2"]])


if __name__ == "__main__":
    unittest.main()
