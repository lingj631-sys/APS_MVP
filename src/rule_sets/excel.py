"""Excel editing surface for rule sets.

The workbook has four sheets: Settings, Groups, Rules, Templates. The Rules sheet
is intentionally flat so a planner can edit it; scope columns (site/line/stage/part)
empty means "all". Import validates every row and reports the Excel cell.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.kernel.registry import TEMPLATES, validate_params
from src.kernel.scope import validate_rule_set
from src.kernel.types import GroupMode, Nature, RuleGroup, RuleInstance, RuleSet, Scope

SETTINGS_HEADERS = ["key", "value"]
GROUP_HEADERS = ["group_id", "priority", "mode"]
RULE_HEADERS = [
    "rule_id",
    "template_id",
    "site",
    "line",
    "stage",
    "part",
    "nature",
    "weight",
    "group_id",
    "priority",
    "inherit",
    "effective_from",
    "effective_to",
    "enabled",
    "params_json",
]
TEMPLATE_HEADERS = ["template_id", "description", "default_nature", "default_group", "params_json"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, headers: List[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for index, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = max(12, min(44, len(header) + 8))


def _example_payload(site: str, rule_set_id: str) -> Dict[str, Any]:
    return {
        "rule_set_id": rule_set_id,
        "site": site,
        "version": 1,
        "groups": [
            {"id": "G0_PHYSICAL", "priority": 0, "mode": "hard"},
            {"id": "G1_DELIVERY", "priority": 10, "mode": "lexicographic"},
            {"id": "G2_CAPACITY", "priority": 20, "mode": "lexicographic"},
            {"id": "G3_CHANGEOVER", "priority": 30, "mode": "weighted"},
        ],
        "rules": [
            {
                "id": f"{site}_CAP_95",
                "template": "capacity_utilization",
                "scope": {},
                "nature": "hard",
                "weight": 100,
                "group": "G2_CAPACITY",
                "priority": 20,
                "params": {"min_util": 0.95, "max_util": 1.0, "tail_exempt": True},
            }
        ],
    }


def _write_settings(ws, site: str, rule_set_id: str, version: int = 1, note: str = "") -> None:
    _style_header(ws, SETTINGS_HEADERS)
    for key, value in (("site", site), ("rule_set_id", rule_set_id), ("version", version), ("note", note)):
        ws.append([key, value])


def _write_groups(ws, groups: List[Dict[str, Any]]) -> None:
    _style_header(ws, GROUP_HEADERS)
    for group in groups:
        ws.append([group.get("id"), group.get("priority", 0), group.get("mode", "weighted")])


def _write_rules(ws, rules: List[Dict[str, Any]]) -> None:
    _style_header(ws, RULE_HEADERS)
    for rule in rules:
        scope = rule.get("scope") or {}
        ws.append(
            [
                rule.get("id"),
                rule.get("template"),
                scope.get("site", ""),
                scope.get("line", ""),
                scope.get("stage", ""),
                scope.get("part", ""),
                rule.get("nature", "soft"),
                rule.get("weight", 1),
                rule.get("group"),
                rule.get("priority", 0),
                rule.get("inherit", "replace"),
                rule.get("effective_from") or "",
                rule.get("effective_to") or "",
                bool(rule.get("enabled", True)),
                json.dumps(rule.get("params") or {}, ensure_ascii=False, sort_keys=True),
            ]
        )


def _write_templates(ws) -> None:
    _style_header(ws, TEMPLATE_HEADERS)
    for template_id in sorted(TEMPLATES):
        spec = TEMPLATES[template_id]
        defaults = {param.name: param.default for param in spec.params}
        ws.append(
            [
                template_id,
                spec.description,
                spec.default_nature,
                spec.default_group,
                json.dumps(defaults, ensure_ascii=False, sort_keys=True),
            ]
        )


def write_template(path: str | Path, site: str = "WL", rule_set_id: Optional[str] = None) -> Path:
    rule_set_id = rule_set_id or f"{site}_RULES"
    payload = _example_payload(site, rule_set_id)
    workbook = Workbook()
    settings = workbook.active
    settings.title = "Settings"
    _write_settings(settings, site, rule_set_id, 1, "编辑后导入；空的作用域列表示全部")
    _write_groups(workbook.create_sheet("Groups"), payload["groups"])
    _write_rules(workbook.create_sheet("Rules"), payload["rules"])
    _write_templates(workbook.create_sheet("Templates"))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def export_rule_set(rule_set: RuleSet, path: str | Path) -> Path:
    payload = rule_set.to_dict()
    workbook = Workbook()
    settings = workbook.active
    settings.title = "Settings"
    _write_settings(settings, rule_set.site, rule_set.rule_set_id, rule_set.version)
    _write_groups(workbook.create_sheet("Groups"), payload["groups"])
    _write_rules(workbook.create_sheet("Rules"), payload["rules"])
    _write_templates(workbook.create_sheet("Templates"))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _rows(ws) -> List[Tuple[int, Dict[str, Tuple[Any, str]]]]:
    headers = {}
    for col in range(1, ws.max_column + 1):
        header = _text(ws.cell(1, col).value)
        if header:
            headers[header] = col
    rows = []
    for row in range(2, ws.max_row + 1):
        values = {}
        empty = True
        for header, col in headers.items():
            cell = ws.cell(row, col)
            values[header] = (cell.value, f"{ws.title}!{cell.coordinate}")
            if _text(cell.value):
                empty = False
        if not empty:
            rows.append((row, values))
    return rows


def _settings(ws) -> Dict[str, str]:
    settings = {}
    for _, values in _rows(ws):
        key = _text(values.get("key", (None, ""))[0])
        value = _text(values.get("value", (None, ""))[0])
        if key:
            settings[key] = value
    return settings


def _bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return _text(value).lower() in ("1", "true", "yes", "y", "是")


def read_workbook(path: str | Path, site: Optional[str] = None) -> Tuple[Optional[Dict[str, Any]], List[Dict[str, str]], List[str]]:
    """Parse and validate an Excel rule set. Errors carry the Excel cell."""
    errors: List[Dict[str, str]] = []
    warnings: List[str] = []
    workbook = load_workbook(path, data_only=True)
    names = {name.lower(): name for name in workbook.sheetnames}
    if "rules" not in names:
        return None, [{"sheet": "Workbook", "cell": "", "rule_id": "", "message": "缺少 Rules 工作表"}], warnings

    settings = _settings(workbook[names["settings"]]) if "settings" in names else {}
    target_site = site or settings.get("site") or "WL"
    rule_set_id = settings.get("rule_set_id") or f"{target_site}_RULES"
    try:
        version = int(settings.get("version") or 1)
    except ValueError:
        version = 1

    groups: List[Dict[str, Any]] = []
    group_ids = set()
    if "groups" in names:
        for _, values in _rows(workbook[names["groups"]]):
            group_id, group_cell = _text(values.get("group_id", (None, ""))[0]), values.get("group_id", (None, ""))[1]
            if not group_id:
                errors.append({"sheet": "Groups", "cell": group_cell, "rule_id": "", "message": "group_id 不能为空"})
                continue
            if group_id in group_ids:
                errors.append({"sheet": "Groups", "cell": group_cell, "rule_id": "", "message": f"group_id 重复: {group_id}"})
                continue
            try:
                priority = int(_text(values.get("priority", (0, ""))[0]) or 0)
            except ValueError:
                errors.append({"sheet": "Groups", "cell": values.get("priority", (None, group_cell))[1], "rule_id": "", "message": "priority 必须是整数"})
                priority = 0
            mode = _text(values.get("mode", ("weighted", ""))[0]) or "weighted"
            try:
                mode_value = GroupMode(mode).value
            except ValueError:
                errors.append({"sheet": "Groups", "cell": values.get("mode", (None, group_cell))[1], "rule_id": "", "message": f"mode 必须是 hard/weighted/lexicographic: {mode}"})
                mode_value = GroupMode.WEIGHTED.value
            group_ids.add(group_id)
            groups.append({"id": group_id, "priority": priority, "mode": mode_value})
    else:
        errors.append({"sheet": "Workbook", "cell": "", "rule_id": "", "message": "缺少 Groups 工作表"})

    rules: List[Dict[str, Any]] = []
    seen_ids = set()
    for _, values in _rows(workbook[names["rules"]]):
        rule_cell = values.get("rule_id", (None, ""))[1]
        rule_id = _text(values.get("rule_id", (None, ""))[0])
        if not rule_id:
            errors.append({"sheet": "Rules", "cell": rule_cell, "rule_id": "", "message": "rule_id 不能为空"})
            continue
        if rule_id in seen_ids:
            errors.append({"sheet": "Rules", "cell": rule_cell, "rule_id": rule_id, "message": "rule_id 重复"})
            continue
        seen_ids.add(rule_id)

        template_id = _text(values.get("template_id", (None, ""))[0])
        template_cell = values.get("template_id", (None, rule_cell))[1]
        if template_id not in TEMPLATES:
            errors.append({"sheet": "Rules", "cell": template_cell, "rule_id": rule_id, "message": f"未知模板: {template_id}"})
            continue

        group_id = _text(values.get("group_id", (None, ""))[0])
        group_cell = values.get("group_id", (None, rule_cell))[1]
        if group_id not in group_ids:
            errors.append({"sheet": "Rules", "cell": group_cell, "rule_id": rule_id, "message": f"分组不存在: {group_id}"})
            continue

        nature_text = _text(values.get("nature", ("soft", ""))[0]) or "soft"
        try:
            nature = Nature(nature_text)
        except ValueError:
            errors.append({"sheet": "Rules", "cell": values.get("nature", (None, rule_cell))[1], "rule_id": rule_id, "message": f"nature 必须是 hard/soft: {nature_text}"})
            continue

        try:
            weight = float(_text(values.get("weight", (1, ""))[0]) or 1)
        except ValueError:
            errors.append({"sheet": "Rules", "cell": values.get("weight", (None, rule_cell))[1], "rule_id": rule_id, "message": "weight 必须是数字"})
            continue
        if not 0 <= weight <= 1000:
            errors.append({"sheet": "Rules", "cell": values.get("weight", (None, rule_cell))[1], "rule_id": rule_id, "message": "weight 必须在 [0,1000]"})
            continue

        try:
            priority = int(_text(values.get("priority", (0, ""))[0]) or 0)
        except ValueError:
            priority = 0
        inherit = _text(values.get("inherit", ("replace", ""))[0]) or "replace"
        if inherit not in ("replace", "add"):
            errors.append({"sheet": "Rules", "cell": values.get("inherit", (None, rule_cell))[1], "rule_id": rule_id, "message": "inherit 必须是 replace/add"})
            continue

        raw_params = values.get("params_json", ("{}", ""))[0]
        try:
            if isinstance(raw_params, dict):
                params = dict(raw_params)
            else:
                params = json.loads(_text(raw_params) or "{}")
            if not isinstance(params, dict):
                raise ValueError("params 不是对象")
            params = validate_params(template_id, params)
        except (ValueError, json.JSONDecodeError) as exc:
            errors.append({"sheet": "Rules", "cell": values.get("params_json", (None, rule_cell))[1], "rule_id": rule_id, "message": f"参数非法: {exc}"})
            continue

        scope = {
            "site": _text(values.get("site", ("", ""))[0]) or "*",
            "line": _text(values.get("line", ("", ""))[0]) or "*",
            "stage": _text(values.get("stage", ("", ""))[0]) or "*",
            "part": _text(values.get("part", ("", ""))[0]) or "*",
        }
        rules.append(
            {
                "id": rule_id,
                "template": template_id,
                "scope": scope,
                "nature": nature.value,
                "weight": weight,
                "group": group_id,
                "priority": priority,
                "params": params,
                "enabled": _bool(values.get("enabled", (True, ""))[0]),
                "inherit": inherit,
                "effective_from": _text(values.get("effective_from", ("", ""))[0]) or None,
                "effective_to": _text(values.get("effective_to", ("", ""))[0]) or None,
            }
        )

    if errors:
        return None, errors, warnings

    payload = {
        "rule_set_id": rule_set_id,
        "site": target_site,
        "version": version,
        "groups": groups,
        "rules": rules,
    }
    try:
        rule_set = RuleSet.from_dict(payload)
        warnings.extend(validate_rule_set(rule_set))
    except (ValueError, KeyError) as exc:
        errors.append({"sheet": "Workbook", "cell": "", "rule_id": "", "message": str(exc)})
        return None, errors, warnings
    return payload, errors, warnings
