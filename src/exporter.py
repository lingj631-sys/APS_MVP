"""
Exporter for the APS MVP.

Writes the schedule into 'Production plan-Done' using the two-row header layout
of the 'Sample' sheet (merged 规格信息 block + one merged date spanning the shift
columns), then appends per-line summary rows: output totals, manufacturing
hours, and shift capacity utilisation with the WE34 95% rule flagged.
"""
from datetime import datetime
from typing import Dict, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.balance_builder import ValidationMatrix, LineLoadRow, MaterialBalanceRow
from src.scheduler import ScheduledTask
from src.shift_manager import MIN_SHIFT_UTILIZATION, SHIFT_SECONDS, ShiftBucket

ShiftKey = Tuple[str, str]

SPEC_LABEL = "规格信息"
BASE_HEADERS = [
    "资源代码", "品目描述", "成品料号", "制造时间", "CT", "TRP",
    "制造开始时间", "备注", "合计",
]
SHIFT_HEADERS = ["早班", "中班", "晚班", "试产"]
OUTPUT_SUMMARY = "产量汇总"
HOURS_SUMMARY = "制造工时汇总(h)"
UTILIZATION_SUMMARY = "班次利用率(%)"

FIRST_SHIFT_COL = len(BASE_HEADERS) + 1
TOTAL_COL = len(BASE_HEADERS)

FILL_TASK_FG = PatternFill("solid", start_color="E8F0FE")
FILL_TASK_SF = PatternFill("solid", start_color="FEF7E0")
FILL_HEADER = PatternFill("solid", start_color="F3F3F3")
FILL_SUMMARY = PatternFill("solid", start_color="EDEDED")
FILL_TAIL = PatternFill("solid", start_color="FFF3CD")
FILL_VIOLATION = PatternFill("solid", start_color="F6C6C6")
THIN = Side(style="thin", color="D3D3D3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COLUMN_WIDTHS = {"A": 13, "B": 30, "C": 14, "D": 10, "E": 6, "F": 6, "G": 17, "H": 24, "I": 9}

BORDER_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
BORDER_MEDIUM_BOTTOM = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="medium", color="94A3B8"),
)


def export_validation_matrix_excel(
    matrix: ValidationMatrix,
    output_excel_path: str,
    site: str = "WL",
) -> None:
    """Export the unified single-sheet Production Plan & Material Balance (PSI) table.

    Exactly matches the on-screen '合并同屏展示 (ALL 37)' view:
    1. Only ONE sheet: '生产排程与物料平衡总表 (PSI)'.
    2. Header with date groups and shift columns.
    3. Section 1: Line Loads Overview (Loading Hours, TRP, Loading Rate).
    4. Section 2: Material Balance Matrix (Customer demand, Production plan, Stock Balance, Coverage days)
       for all Finished Goods (FG) and Semi-Finished Goods (SF).
    5. Froze panes at E3 for seamless horizontal and vertical scrolling.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "生产排程与物料平衡总表 (PSI)"
    try:
        ws.views.sheetView[0].showGridLines = True
    except Exception:
        pass

    shift_cols = matrix.shift_columns
    total_cols = 4 + len(shift_cols)

    # 1. Row 1 & 2: Header Styling & Content
    fill_hdr = PatternFill("solid", start_color="F1F5F9")
    font_hdr = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
    font_sub_hdr = Font(name="Segoe UI", size=9, bold=True, color="475569")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20

    # Base headers (Columns 1-4)
    base_headers_r1 = ["Site / 料号", "产线 / 品名描述", "监控指标 / 业务项", "定额 / 基准参数"]
    base_headers_r2 = ["站点 / Part No", "Line / Description", "Metric / Action", "Standard / Ref"]
    for c_idx in range(1, 5):
        ws.cell(1, c_idx).value = base_headers_r1[c_idx - 1]
        ws.cell(2, c_idx).value = base_headers_r2[c_idx - 1]

    # Shift headers (Columns 5..total_cols)
    # Group consecutive columns with the exact same date in Row 1
    c_start = 5
    while c_start <= total_cols:
        d_val = shift_cols[c_start - 5]["date"]
        c_end = c_start
        while (c_end + 1 <= total_cols) and (shift_cols[c_end - 5 + 1]["date"] == d_val):
            c_end += 1
        if c_end > c_start:
            ws.merge_cells(start_row=1, start_column=c_start, end_row=1, end_column=c_end)
        ws.cell(1, c_start).value = d_val
        c_start = c_end + 1

    for idx, col_meta in enumerate(shift_cols, start=5):
        ws.cell(2, idx).value = col_meta["shift"]

    for r in (1, 2):
        for c in range(1, total_cols + 1):
            cell = ws.cell(r, c)
            cell.font = font_hdr if r == 1 else font_sub_hdr
            cell.fill = fill_hdr
            cell.border = BORDER_THIN
            cell.alignment = align_center

    # 2. Section 1 Banner: Line Load Overview
    curr_row = 3
    ws.row_dimensions[curr_row].height = 24
    ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=total_cols)
    banner_ll = ws.cell(curr_row, 1)
    banner_ll.value = "【上表】产线负荷概览 (Line Load Overview - 纯工时 / TRP / 班次利用率)    定额: 单班8h · 门槛: ≥95% (尾班可豁免)"
    banner_ll.font = Font(name="Segoe UI", size=10, bold=True, color="312E81")
    banner_ll.fill = PatternFill("solid", start_color="EEF2FF")
    banner_ll.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(1, total_cols + 1):
        ws.cell(curr_row, c).fill = PatternFill("solid", start_color="EEF2FF")
        ws.cell(curr_row, c).border = BORDER_THIN

    # Line Loads Rows
    for row_data in matrix.line_loads:
        curr_row += 1
        ws.row_dimensions[curr_row].height = 19
        ws.cell(curr_row, 1).value = row_data.site
        ws.cell(curr_row, 1).font = Font(name="Segoe UI", size=9, color="475569")
        ws.cell(curr_row, 1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 1).border = BORDER_THIN

        ws.cell(curr_row, 2).value = row_data.line
        ws.cell(curr_row, 2).font = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
        ws.cell(curr_row, 2).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 2).border = BORDER_THIN

        ws.cell(curr_row, 3).value = row_data.metric
        ws.cell(curr_row, 3).font = Font(
            name="Segoe UI",
            size=9,
            bold=(row_data.metric == "Loading Rate"),
            color="3730A3" if row_data.metric == "Loading Rate" else "334155",
        )
        ws.cell(curr_row, 3).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 3).border = BORDER_THIN

        # Col 4 Standard / Ref
        if row_data.metric == "Loading Rate":
            ref_txt = "门槛 ≥95%"
        elif row_data.metric == "TRP":
            ref_txt = f"标准: {0.71 if row_data.line == 'PGLINKAGE' else 0.85}"
        else:
            ref_txt = "单班定额: 12.0h" if site == "SY38" else "单班定额: 8.0h"
        ws.cell(curr_row, 4).value = ref_txt
        ws.cell(curr_row, 4).font = Font(
            name="Segoe UI",
            size=9,
            bold=(row_data.metric == "Loading Rate"),
            color="065F46" if row_data.metric == "Loading Rate" else "64748B",
        )
        ws.cell(curr_row, 4).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(curr_row, 4).border = BORDER_THIN

        # Shift values
        for idx, col_meta in enumerate(shift_cols, start=5):
            val = row_data.shift_values.get(col_meta["key"], 0)
            c = ws.cell(curr_row, idx)
            c.alignment = align_center
            c.border = BORDER_THIN
            if row_data.metric == "Loading Rate":
                if isinstance(val, (int, float)) and val > 0:
                    c.value = f"{val:.1f}%"
                    if val >= 95.0:
                        c.fill = PatternFill("solid", start_color="D1FAE5")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color="065F46")
                    else:
                        c.fill = PatternFill("solid", start_color="FEE2E2")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")
                else:
                    c.value = "-"
                    c.font = Font(name="Segoe UI", size=9, color="94A3B8")
            elif row_data.metric == "Loading Hours":
                if isinstance(val, (int, float)) and val > 0:
                    c.value = round(val, 2)
                    c.font = Font(name="Segoe UI", size=9, color="1E293B")
                else:
                    c.value = "-"
                    c.font = Font(name="Segoe UI", size=9, color="94A3B8")
            else:  # TRP
                c.value = val
                c.font = Font(name="Segoe UI", size=9, color="64748B")

        if row_data.metric == "Loading Rate":
            for col_idx in range(1, total_cols + 1):
                ws.cell(curr_row, col_idx).border = BORDER_MEDIUM_BOTTOM

    # 3. Section 2 Banner: Material Balance Matrix
    curr_row += 1
    ws.row_dimensions[curr_row].height = 24
    ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=total_cols)
    banner_mat = ws.cell(curr_row, 1)
    banner_mat.value = (
        f"【下表】物料动态平衡表 (PSI Balance Matrix - 需求 / 排产 / 结存 / 覆盖天数)    "
        f"共 {len(matrix.fg_balance) + len(matrix.sf_balance)} 种物料 · 产成品 (FG) + 半成品 (SFIN) 合并同屏展示"
    )
    banner_mat.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
    banner_mat.fill = PatternFill("solid", start_color="D1FAE5")
    banner_mat.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(1, total_cols + 1):
        ws.cell(curr_row, c).fill = PatternFill("solid", start_color="D1FAE5")
        ws.cell(curr_row, c).border = BORDER_THIN

    # Sub-header row for Material Balance
    curr_row += 1
    ws.row_dimensions[curr_row].height = 19
    sub_headers = ["料号 (PN)", "品名与主参数", "指标项 (Des)", "基准参数 (Reference)"]
    for col_idx, h in enumerate(sub_headers, start=1):
        c = ws.cell(curr_row, col_idx)
        c.value = h
        c.font = Font(name="Segoe UI", size=9, bold=True, color="475569")
        c.fill = PatternFill("solid", start_color="F8FAFC")
        c.border = BORDER_THIN
        c.alignment = Alignment(horizontal="right" if col_idx == 4 else "left", vertical="center")

    for idx, col_meta in enumerate(shift_cols, start=5):
        c = ws.cell(curr_row, idx)
        c.value = f"{col_meta['date']} {col_meta['shift']}"
        c.font = Font(name="Segoe UI", size=8, bold=True, color="64748B")
        c.fill = PatternFill("solid", start_color="F8FAFC")
        c.border = BORDER_THIN
        c.alignment = align_center

    # Loop all materials (FG + SF = 37 parts)
    all_materials = list(matrix.fg_balance) + list(matrix.sf_balance)
    for mat in all_materials:
        # 1. Customer demand
        curr_row += 1
        ws.row_dimensions[curr_row].height = 19
        ws.cell(curr_row, 1).value = mat.part_no
        ws.cell(curr_row, 1).font = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
        ws.cell(curr_row, 1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 1).border = BORDER_THIN

        ws.cell(curr_row, 2).value = mat.description
        ws.cell(curr_row, 2).font = Font(name="Segoe UI", size=9, color="1E293B")
        ws.cell(curr_row, 2).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 2).border = BORDER_THIN

        ws.cell(curr_row, 3).value = "Customer demand"
        ws.cell(curr_row, 3).font = Font(name="Segoe UI", size=9, color="475569")
        ws.cell(curr_row, 3).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 3).border = BORDER_THIN

        ws.cell(curr_row, 4).value = f"总需求: {int(mat.total_demand)} 件"
        ws.cell(curr_row, 4).font = Font(name="Segoe UI", size=9, color="475569")
        ws.cell(curr_row, 4).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(curr_row, 4).border = BORDER_THIN

        for idx, col_meta in enumerate(shift_cols, start=5):
            d_val = mat.demand_by_shift.get(col_meta["key"], 0)
            c = ws.cell(curr_row, idx)
            c.alignment = align_center
            c.border = BORDER_THIN
            if d_val > 0:
                c.value = int(d_val)
                c.fill = PatternFill("solid", start_color="FEF3C7")
                c.font = Font(name="Segoe UI", size=9, bold=True, color="92400E")
            else:
                c.value = "-"
                c.font = Font(name="Segoe UI", size=9, color="CBD5E1")

        # 2. Production plan
        curr_row += 1
        ws.row_dimensions[curr_row].height = 19
        ws.cell(curr_row, 1).value = f"{mat.line} · {mat.site}"
        ws.cell(curr_row, 1).font = Font(name="Segoe UI", size=8, color="64748B")
        ws.cell(curr_row, 1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 1).border = BORDER_THIN

        ws.cell(curr_row, 2).value = f"MPQ: {int(mat.mpq)} · TRP: {mat.trp}"
        ws.cell(curr_row, 2).font = Font(name="Segoe UI", size=8, color="64748B")
        ws.cell(curr_row, 2).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 2).border = BORDER_THIN

        ws.cell(curr_row, 3).value = "Production plan"
        ws.cell(curr_row, 3).font = Font(name="Segoe UI", size=9, bold=True, color="1D4ED8")
        ws.cell(curr_row, 3).fill = PatternFill("solid", start_color="EFF6FF")
        ws.cell(curr_row, 3).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 3).border = BORDER_THIN

        ws.cell(curr_row, 4).value = f"节拍 CT: {mat.ct_seconds} 秒"
        ws.cell(curr_row, 4).font = Font(name="Segoe UI", size=9, color="1D4ED8")
        ws.cell(curr_row, 4).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(curr_row, 4).border = BORDER_THIN

        for idx, col_meta in enumerate(shift_cols, start=5):
            p_val = mat.plan_by_shift.get(col_meta["key"], 0)
            c = ws.cell(curr_row, idx)
            c.alignment = align_center
            c.border = BORDER_THIN
            if p_val > 0:
                c.value = int(p_val)
                c.fill = PatternFill("solid", start_color="DBEAFE")
                c.font = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
            else:
                c.value = "-"
                c.font = Font(name="Segoe UI", size=9, color="CBD5E1")

        # 3. Stock Balance
        curr_row += 1
        ws.row_dimensions[curr_row].height = 19
        ws.cell(curr_row, 1).value = ""
        ws.cell(curr_row, 1).border = BORDER_THIN
        ws.cell(curr_row, 2).value = ""
        ws.cell(curr_row, 2).border = BORDER_THIN
        ws.cell(curr_row, 3).value = "Stock Balance"
        ws.cell(curr_row, 3).font = Font(name="Segoe UI", size=9, color="1E293B")
        ws.cell(curr_row, 3).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(curr_row, 3).border = BORDER_THIN

        ws.cell(curr_row, 4).value = f"期初: {int(mat.initial_stock)} 件"
        ws.cell(curr_row, 4).font = Font(name="Segoe UI", size=9, color="047857")
        ws.cell(curr_row, 4).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(curr_row, 4).border = BORDER_THIN

        for idx, col_meta in enumerate(shift_cols, start=5):
            s_val = mat.stock_by_shift.get(col_meta["key"], 0)
            c = ws.cell(curr_row, idx)
            c.value = int(s_val)
            c.alignment = align_center
            c.border = BORDER_THIN
            if s_val < 0:
                c.fill = PatternFill("solid", start_color="FEE2E2")
                c.font = Font(name="Segoe UI", size=9, bold=True, color="B91C1C")
            else:
                c.font = Font(name="Segoe UI", size=9, color="1E293B")

        # 4. Coverage days
        curr_row += 1
        ws.row_dimensions[curr_row].height = 19
        ws.cell(curr_row, 1).value = ""
        ws.cell(curr_row, 2).value = ""
        ws.cell(curr_row, 3).value = "Coverage days"
        ws.cell(curr_row, 3).font = Font(name="Segoe UI", size=9, color="64748B")
        ws.cell(curr_row, 3).alignment = Alignment(horizontal="left", vertical="center")

        ws.cell(curr_row, 4).value = f"日均: {mat.avg_daily_demand:.1f} 件/天"
        ws.cell(curr_row, 4).font = Font(name="Segoe UI", size=9, color="64748B")
        ws.cell(curr_row, 4).alignment = Alignment(horizontal="right", vertical="center")

        for idx, col_meta in enumerate(shift_cols, start=5):
            c_val = mat.coverage_by_shift.get(col_meta["key"], 999.0)
            c = ws.cell(curr_row, idx)
            c.alignment = align_center
            if c_val >= 999 or c_val == 999.0:
                c.value = 999
                c.font = Font(name="Segoe UI", size=9, color="94A3B8")
            else:
                c.value = round(c_val, 1)
                if c_val < 2.0:
                    c.fill = PatternFill("solid", start_color="FEF2F2")
                    c.font = Font(name="Segoe UI", size=9, bold=True, color="DC2626")
                else:
                    c.font = Font(name="Segoe UI", size=9, color="475569")

        # Separator bottom border on the 4th row
        for col_idx in range(1, total_cols + 1):
            ws.cell(curr_row, col_idx).border = BORDER_MEDIUM_BOTTOM

    # Freeze panes at E3: keeps Row 1-2 headers and Col 1-4 columns frozen
    ws.freeze_panes = "E3"

    # Set column widths
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 22
    for col_idx in range(5, total_cols + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 9.5

    wb.save(output_excel_path)
    print(f"Successfully exported single-sheet validation matrix to {output_excel_path}")


def export_production_plan(
    template_excel_path: str,
    output_excel_path: str,
    scheduled_tasks: List[ScheduledTask],
    shifts: List[ShiftBucket],
) -> None:
    dates = _ordered_dates(shifts)
    total_cols = TOTAL_COL + len(SHIFT_HEADERS) * len(dates)
    col_of = _shift_columns(dates)

    wb = openpyxl.load_workbook(template_excel_path, data_only=False)
    ws = (
        wb["Production plan-Done"]
        if "Production plan-Done" in wb.sheetnames
        else wb.create_sheet("Production plan-Done")
    )
    _reset_sheet(ws)
    _write_header(ws, dates)

    row = 3
    for line in _ordered_lines(scheduled_tasks):
        line_tasks = sorted(
            (t for t in scheduled_tasks if t.line == line),
            key=lambda t: (t.start_dt, t.part_no),
        )
        fill = FILL_TASK_FG if _is_fg_line(line_tasks) else FILL_TASK_SF
        for task in line_tasks:
            _write_task_row(ws, row, task, col_of, fill)
            row += 1
        row = _write_summary_rows(ws, row, line, line_tasks, col_of) + 1

    _write_grand_total(ws, row, scheduled_tasks, col_of)
    _finish_sheet(ws, total_cols, row)

    wb.save(output_excel_path)
    print(f"Successfully exported {len(scheduled_tasks)} scheduled tasks to {output_excel_path}")


def _ordered_dates(shifts: List[ShiftBucket]) -> List[str]:
    seen: List[str] = []
    for shift in shifts:
        if shift.date_str not in seen:
            seen.append(shift.date_str)
    return seen


def _shift_columns(dates: List[str]) -> Dict[ShiftKey, int]:
    col_of: Dict[ShiftKey, int] = {}
    for i, date_str in enumerate(dates):
        for j, shift_name in enumerate(SHIFT_HEADERS[:3]):
            col_of[(date_str, shift_name)] = FIRST_SHIFT_COL + i * len(SHIFT_HEADERS) + j
    return col_of


def _ordered_lines(tasks: List[ScheduledTask]) -> List[str]:
    seen: List[str] = []
    for task in sorted(tasks, key=lambda t: t.start_dt):
        if task.line not in seen:
            seen.append(task.line)
    return seen


def _is_fg_line(tasks: List[ScheduledTask]) -> bool:
    return any(t.stage == "FG" for t in tasks)


def _reset_sheet(ws) -> None:
    for merge in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merge))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)


def _write_header(ws, dates: List[str]) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=TOTAL_COL)
    ws.cell(1, 3).value = SPEC_LABEL
    for i, header in enumerate(BASE_HEADERS, start=1):
        ws.cell(2, i).value = header

    for i, date_str in enumerate(dates):
        first_col = FIRST_SHIFT_COL + i * len(SHIFT_HEADERS)
        ws.merge_cells(start_row=1, start_column=first_col, end_row=1, end_column=first_col + 3)
        cell = ws.cell(1, first_col)
        cell.value = _parse_date(date_str)
        cell.number_format = "M/D"
        for j, shift_name in enumerate(SHIFT_HEADERS):
            ws.cell(2, first_col + j).value = shift_name


def _parse_date(date_str: str) -> datetime:
    month, day = (int(part) for part in date_str.split("/"))
    return datetime(2026, month, day)


def _write_task_row(ws, row: int, task: ScheduledTask, col_of: Dict[ShiftKey, int], fill: PatternFill) -> None:
    ws.cell(row, 1).value = task.line
    ws.cell(row, 2).value = task.description
    ws.cell(row, 3).value = task.part_no
    ws.cell(row, 4).value = f"{task.duration_hours:.1f}H"
    ws.cell(row, 5).value = task.ct_seconds
    ws.cell(row, 6).value = task.trp
    ws.cell(row, 7).value = task.start_dt
    ws.cell(row, 7).number_format = "yyyy-mm-dd hh:mm"
    ws.cell(row, 8).value = task.order_id
    ws.cell(row, 9).value = int(round(task.quantity))
    for key, qty in task.shift_breakdown.items():
        col = col_of.get(key)
        if col is not None and qty:
            ws.cell(row, col).value = int(qty)
    for col in range(1, TOTAL_COL + 1):
        ws.cell(row, col).fill = fill


def _shift_loads(tasks: List[ScheduledTask]) -> Dict[ShiftKey, float]:
    loads: Dict[ShiftKey, float] = {}
    for task in tasks:
        for key, seconds in task.shift_seconds.items():
            loads[key] = loads.get(key, 0.0) + seconds
    return loads


def _write_summary_rows(ws, row: int, line: str, tasks: List[ScheduledTask], col_of: Dict[ShiftKey, int]) -> int:
    loads = _shift_loads(tasks)
    worked = sorted(loads, key=_shift_sort_key)
    closing = worked[-1] if worked else None
    output_row, hours_row, util_row = row, row + 1, row + 2

    for target, label in ((output_row, OUTPUT_SUMMARY), (hours_row, HOURS_SUMMARY), (util_row, UTILIZATION_SUMMARY)):
        ws.cell(target, 1).value = line
        ws.cell(target, 2).value = label

    ws.cell(output_row, 9).value = sum(int(round(t.quantity)) for t in tasks)
    ws.cell(hours_row, 4).value = f"{sum(t.duration_hours for t in tasks):.1f}H"
    ws.cell(hours_row, 9).value = round(sum(loads.values()) / 3600.0, 2)
    if worked:
        ws.cell(util_row, 9).value = round(
            sum(loads.values()) / (len(worked) * SHIFT_SECONDS) * 100, 1
        )

    for key, seconds in loads.items():
        col = col_of.get(key)
        if col is None:
            continue
        ws.cell(output_row, col).value = sum(int(t.shift_breakdown.get(key, 0)) for t in tasks)
        ws.cell(hours_row, col).value = round(seconds / 3600.0, 2)
        utilization = seconds / SHIFT_SECONDS
        cell = ws.cell(util_row, col)
        cell.value = round(utilization * 100, 1)
        cell.number_format = "0.0"
        if utilization < MIN_SHIFT_UTILIZATION - 1e-9:
            cell.fill = FILL_TAIL if key == closing else FILL_VIOLATION

    return util_row


def _shift_sort_key(key: ShiftKey) -> Tuple[int, int, int]:
    month, day = (int(part) for part in key[0].split("/"))
    return month, day, SHIFT_HEADERS.index(key[1]) if key[1] in SHIFT_HEADERS else 9


def _write_grand_total(ws, row: int, tasks: List[ScheduledTask], col_of: Dict[ShiftKey, int]) -> None:
    """Quantity totals, matching the total row of the 'Sample' sheet.

    Per-shift hours are deliberately not summed here: two lines share a shift
    slot, so an hours total would read as a capacity breach.
    """
    outputs: Dict[ShiftKey, int] = {}
    for task in tasks:
        for key, qty in task.shift_breakdown.items():
            outputs[key] = outputs.get(key, 0) + int(qty)

    ws.cell(row, 1).value = "全厂"
    ws.cell(row, 2).value = "总计(数量)"
    ws.cell(row, 4).value = f"{sum(t.duration_hours for t in tasks):.1f}H"
    ws.cell(row, 9).value = sum(int(round(t.quantity)) for t in tasks)
    for key, qty in outputs.items():
        col = col_of.get(key)
        if col is not None:
            ws.cell(row, col).value = qty
    for col in range(1, TOTAL_COL + 1):
        ws.cell(row, col).fill = FILL_SUMMARY


def _finish_sheet(ws, total_cols: int, last_row: int) -> None:
    for row in (1, 2):
        for col in range(1, total_cols + 1):
            cell = ws.cell(row, col)
            cell.font = Font(bold=True)
            cell.fill = FILL_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
    for row in range(3, last_row + 1):
        for col in range(1, total_cols + 1):
            ws.cell(row, col).border = BORDER
    for letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    for col in range(FIRST_SHIFT_COL, total_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 7
    ws.freeze_panes = f"{get_column_letter(FIRST_SHIFT_COL)}3"
