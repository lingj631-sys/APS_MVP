"""Application service for auditing manually adjusted Excel schedules.

Extracted from src/server.py to keep the HTTP router thin and maintain a clean
separation between transport and scheduling reconstruction logic.
"""
import io
from collections import defaultdict
from datetime import datetime, timedelta
from typing import Any, Dict, List

import openpyxl
from fastapi import HTTPException

from src.audit import (
    AuditConfig,
    PlanRow,
    ShiftInfo,
    audit_schedule_part_level,
)
from src.data_loader import (
    ProductionInputSnapshot,
    build_requirement_plan,
)
from src.balance_builder import build_validation_matrix
from src.metric_labels import labels_for
from src.rules import get_rules_for_site
from src.scheduler import APSScheduler, ScheduleResult, ScheduledTask
from src.shift_manager import CalendarMode, DEFAULT_CHANGEOVER_MINUTES
from src.validator import (
    diagnose_empty_shifts,
    validate_schedule,
)


def process_manual_excel_audit(
    content: bytes,
    site: str,
    session_state: Dict[str, Any],
    site_settings: Dict[str, Any],
    filename: str = "manual.xlsx",
) -> Dict[str, Any]:
    """Parse a user-uploaded Excel schedule (PSI or plan-Done), reconstruct

    scheduled tasks, and execute both validator and independent audit checks.
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"无法打开上传的 Excel 文件: {e}")

    site_upper = (site or "WL").upper()
    snap: ProductionInputSnapshot = session_state.get("snapshot")
    if snap is None:
        from src.data_loader import load_aps_database
        snap = load_aps_database("dataSource.xlsx", site_filter=site_upper)

    settings = site_settings
    horizon_days = settings.get("horizon_days", 10)
    start_date_str = settings.get("start_date", "2026-10-12")
    cal_mode_str = settings.get("calendar_mode", "SEVEN_ZERO")
    y, m, d = [int(p) for p in start_date_str.split("-")]
    start_dt = datetime(y, m, d)

    try:
        cal_mode = CalendarMode(cal_mode_str)
    except Exception:
        cal_mode = CalendarMode.SEVEN_ZERO

    plan = session_state.get("plan") or build_requirement_plan(snap)
    scheduler = APSScheduler(
        plan=plan,
        start_datetime=start_dt,
        planning_days=horizon_days,
        calendar_mode=cal_mode,
        site=site_upper,
    )
    shifts = scheduler.shifts
    index_of: Dict[Any, int] = {}
    shift_by_index: Dict[int, Any] = {}
    for s in shifts:
        index_of[(s.date_str, s.shift_name)] = s.index
        index_of[f"{s.date_str} {s.shift_name}"] = s.index
        shift_by_index[s.index] = s

    # Detect worksheet
    ws = None
    for name in ["生产排程与物料平衡总表 (PSI)", "Production plan-Done"]:
        if name in wb.sheetnames:
            ws = wb[name]
            break
    if ws is None:
        ws = wb.active

    tasks_raw: List[Dict[str, Any]] = []
    if "生产排程" in ws.title or "PSI" in ws.title:
        shift_cols: Dict[int, str] = {}
        last_date = ""
        for col in range(5, ws.max_column + 1):
            d_val = ws.cell(1, col).value
            s_val = ws.cell(2, col).value
            if d_val:
                last_date = d_val if isinstance(d_val, str) else f"{d_val.month}/{d_val.day}"
            if s_val:
                s_name = str(s_val).strip()
                if s_name == "早":
                    s_name = "早班"
                elif s_name == "中":
                    s_name = "中班"
                elif s_name == "晚":
                    s_name = "晚班"
                shift_cols[col] = f"{last_date} {s_name}"
        for r in range(3, ws.max_row + 1):
            metric = str(ws.cell(r, 3).value or "").strip()
            if "Production plan" in metric or "排产计划" in metric:
                raw_c1 = str(ws.cell(r, 1).value or "").strip()
                raw_pn = str(ws.cell(r - 1, 1).value or "").strip()
                pn = raw_pn if raw_pn and not raw_pn.startswith("Site") and not raw_pn.startswith("料号") and not raw_pn.startswith("站点") else raw_c1
                line = raw_c1.split("·")[0].strip() if "·" in raw_c1 else raw_c1

                param = snap.fg_params.get(pn) or snap.sf_params.get(pn)
                if param:
                    line = param.line

                breakdown: Dict[str, float] = {}
                for c, sk in shift_cols.items():
                    v = ws.cell(r, c).value
                    if isinstance(v, (int, float)) and v > 0:
                        breakdown[sk] = float(v)
                if breakdown and pn:
                    tasks_raw.append({"part_no": pn, "line": line, "breakdown": breakdown})
    else:
        shift_cols = {}
        last_date = ""
        for col in range(10, ws.max_column + 1):
            d_val = ws.cell(1, col).value
            s_val = ws.cell(2, col).value
            if d_val:
                if isinstance(d_val, datetime):
                    last_date = f"{d_val.month}/{d_val.day}"
                else:
                    last_date = str(d_val).strip()
            if s_val:
                s_name = str(s_val).strip()
                if s_name == "早":
                    s_name = "早班"
                elif s_name == "中":
                    s_name = "中班"
                elif s_name == "晚":
                    s_name = "晚班"
                shift_cols[col] = f"{last_date} {s_name}"
        for r in range(3, ws.max_row + 1):
            line = str(ws.cell(r, 1).value or "").strip()
            pn = str(ws.cell(r, 3).value or "").strip()
            if not pn:
                continue
            breakdown = {}
            for c, sk in shift_cols.items():
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)) and v > 0:
                    breakdown[sk] = float(v)
            if breakdown:
                tasks_raw.append({"part_no": pn, "line": line, "breakdown": breakdown})

    if not tasks_raw:
        raise HTTPException(
            status_code=400,
            detail="未能从上传的 Excel 文件中解析出排产计划数据，请确认文件包含【生产排程与物料平衡总表 (PSI)】或【Production plan-Done】工作表且包含分班排产数值。"
        )

    # Reconstruct ScheduledTask objects
    manual_tasks: List[ScheduledTask] = []
    task_idx = 1
    for item in tasks_raw:
        pn = item["part_no"]
        line = item["line"]
        breakdown = item["breakdown"]

        is_fg = pn in snap.fg_params
        param = snap.fg_params.get(pn) if is_fg else snap.sf_params.get(pn)
        stage = "FG" if is_fg else "SF"
        default_line = "PGLINKAGE" if is_fg else "轴套6"
        task_line = param.line if param else (line if (line and not line.startswith("MPQ") and line != "None") else default_line)
        ct = param.ct_seconds if param else (40.0 if is_fg else 25.0)
        trp = param.trp if param else (0.71 if is_fg else 0.85)
        mpq = int(param.mpq) if param else 1
        desc = param.description if param else pn

        mapped_breakdown = {}
        for sk, qty in breakdown.items():
            if isinstance(sk, tuple):
                mapped_breakdown[sk] = int(round(qty))
            else:
                parts = str(sk).split()
                tuple_key = (parts[0], parts[1]) if len(parts) >= 2 else (str(sk), "早班")
                mapped_breakdown[tuple_key] = int(round(qty))

        matching = sorted([t for t in plan.tasks if t.part_no == pn], key=lambda t: t.due_datetime or datetime.max)
        if not matching or len(matching) == 1:
            matched_t = matching[0] if matching else None
            tot_qty = sum(mapped_breakdown.values())
            active_indices = [index_of[k] for k in mapped_breakdown if k in index_of]
            if not active_indices:
                continue
            min_s, max_s = min(active_indices), max(active_indices)
            shift_secs = {k: v * (ct / trp) for k, v in mapped_breakdown.items()}
            manual_tasks.append(ScheduledTask(
                task_id=matched_t.task_id if matched_t else f"MANUAL_{stage}_{pn}_{task_idx:02d}",
                stage=stage,
                line=task_line,
                part_no=pn,
                description=desc,
                quantity=float(tot_qty),
                ct_seconds=ct,
                trp=trp,
                mpq=mpq,
                lot_size=mpq if is_fg else 1,
                lots=int(tot_qty // mpq) if is_fg else int(tot_qty),
                order_id=matched_t.order_id if matched_t else f"ORD_MANUAL_{pn}",
                due_datetime=matched_t.due_datetime if matched_t else (shift_by_index[min_s].start_dt + timedelta(days=7)),
                start_dt=shift_by_index[min_s].start_dt,
                end_dt=shift_by_index[max_s].end_dt,
                duration_hours=sum(shift_secs.values()) / 3600.0,
                shift_breakdown=mapped_breakdown,
                shift_seconds=shift_secs,
            ))
            task_idx += 1
        else:
            rem_needs = {t.task_id: t.quantity for t in matching}
            task_breakdowns = defaultdict(dict)
            sorted_shifts = sorted(mapped_breakdown.keys(), key=lambda k: index_of.get(k, 0))
            for sk in sorted_shifts:
                avail = mapped_breakdown[sk]
                for t in matching:
                    if avail <= 0:
                        break
                    rem = rem_needs[t.task_id]
                    if rem <= 0:
                        continue
                    take = min(avail, rem)
                    task_breakdowns[t.task_id][sk] = take
                    rem_needs[t.task_id] -= take
                    avail -= take
            if avail > 0:
                last_t = matching[-1]
                last_sk = sorted_shifts[-1]
                task_breakdowns[last_t.task_id][last_sk] = task_breakdowns[last_t.task_id].get(last_sk, 0) + avail

            for t in matching:
                tb = task_breakdowns.get(t.task_id, {})
                tot_qty = sum(tb.values())
                if tot_qty <= 0:
                    continue
                active_indices = [index_of[k] for k in tb if k in index_of]
                if not active_indices:
                    continue
                min_s, max_s = min(active_indices), max(active_indices)
                shift_secs = {k: v * (ct / trp) for k, v in tb.items()}
                manual_tasks.append(ScheduledTask(
                    task_id=t.task_id,
                    stage=stage,
                    line=task_line,
                    part_no=pn,
                    description=desc,
                    quantity=float(tot_qty),
                    ct_seconds=ct,
                    trp=trp,
                    mpq=mpq,
                    lot_size=mpq if is_fg else 1,
                    lots=int(tot_qty // mpq) if is_fg else int(tot_qty),
                    order_id=t.order_id,
                    due_datetime=t.due_datetime,
                    start_dt=shift_by_index[min_s].start_dt,
                    end_dt=shift_by_index[max_s].end_dt,
                    duration_hours=sum(shift_secs.values()) / 3600.0,
                    shift_breakdown=tb,
                    shift_seconds=shift_secs,
                ))

    manual_result = ScheduleResult(
        tasks=manual_tasks,
        shifts=shifts,
        strict_utilization=True,
        status_name="MANUAL_AUDIT",
        objective_value=0.0,
        wall_time_seconds=0.02,
    )

    baseline_result = session_state.get("result")
    if baseline_result and getattr(baseline_result, "tasks", None):
        baseline_task_map = {t.task_id: t for t in baseline_result.tasks}
        for mt in manual_tasks:
            bt = baseline_task_map.get(mt.task_id)
            if bt:
                bt_min_s = min([index_of[k] for k in bt.shift_breakdown if k in index_of], default=999)
                if bt_min_s < 3:
                    setattr(mt, "is_locked", True)

    allow_105_fg = settings.get("allow_105_utilization_fg", True)
    allow_105_sf = settings.get("allow_105_utilization_sf", True)
    checks = validate_schedule(
        plan,
        manual_result,
        site=site_upper,
        allow_105_fg=allow_105_fg,
        allow_105_sf=allow_105_sf,
        baseline_result=baseline_result,
    )

    audit_rows = []
    for item in tasks_raw:
        pn = item["part_no"]
        param = snap.fg_params.get(pn) or snap.sf_params.get(pn)
        stage = "FG" if pn in snap.fg_params else "SF"
        line = param.line if param else item["line"]
        for shift_key, qty in item["breakdown"].items():
            audit_rows.append(
                PlanRow(
                    task_id=f"MANUAL::{pn}",
                    stage=stage,
                    line=line,
                    part_no=pn,
                    order_id="",
                    shift_key=str(shift_key),
                    quantity=float(qty),
                )
            )
    shift_infos = [
        ShiftInfo(
            key=f"{s.date_str} {s.shift_name}",
            index=s.index,
            date_str=s.date_str,
            shift_name=s.shift_name,
            start_dt=s.start_dt,
            end_dt=s.end_dt,
            duration_hours=s.duration_hours,
        )
        for s in shifts
    ]
    independent_audit = audit_schedule_part_level(
        snap,
        audit_rows,
        shift_infos,
        AuditConfig(
            setup_minutes=DEFAULT_CHANGEOVER_MINUTES,
            allow_105=allow_105_fg or allow_105_sf,
            delivery_tolerance_minutes=24.0 if (allow_105_fg or allow_105_sf) else 0.0,
        ),
    )
    empty_shift_reasons = diagnose_empty_shifts(manual_tasks, shifts, plan)
    matrix = build_validation_matrix(snap, plan, manual_result, site=site_upper)

    site_rules = get_rules_for_site(site_upper)
    check_status: Dict[str, bool] = {}
    check_by_rule: Dict[str, Any] = {}
    for c in checks:
        for r_id in c.rule_ids:
            check_status[r_id] = check_status.get(r_id, True) and c.passed
            check_by_rule[r_id] = c

    compliance_rows = []
    for r_id, rule in sorted(site_rules.items(), key=lambda x: x[1].row_index):
        if r_id in check_status:
            eff_status = "PASS" if check_status[r_id] else "FAIL"
        else:
            eff_status = rule.status.value

        matched_check = check_by_rule.get(r_id)
        compliance_rows.append({
            "rule_id": rule.rule_id,
            "row_index": rule.row_index,
            "stage": rule.stage,
            "category": rule.category.value,
            "nature": rule.nature.value,
            "name": rule.name,
            "status": eff_status,
            "weight": rule.weight,
            "note": rule.validation_check or rule.landing_layer,
            "ai_logic": rule.ai_logic,
            "is_checked": True,
            "metrics": getattr(matched_check, "metrics", {}) if matched_check else {},
            "metric_labels": labels_for(getattr(matched_check, "metrics", {}) if matched_check else {}),
            "calculation_details": getattr(matched_check, "calculation_details", []) if matched_check else [],
            "table_data": getattr(matched_check, "table_data", None) if matched_check else None,
            "violations": getattr(matched_check, "violations", []) if matched_check else [],
        })

    checks_data = [
        {
            "name": c.name,
            "passed": c.passed,
            "rule_ids": c.rule_ids,
            "violations": c.violations,
            "total_violations": len(c.violations),
            "metrics": getattr(c, "metrics", {}),
            "metric_labels": labels_for(getattr(c, "metrics", {})),
            "calculation_details": getattr(c, "calculation_details", []),
            "table_data": getattr(c, "table_data", None),
        }
        for c in checks
    ]

    line_loads_data = [
        {"site": row.site, "line": row.line, "metric": row.metric, "shift_values": row.shift_values}
        for row in matrix.line_loads
    ]

    fg_balance_data = [
        {
            "site": row.site, "line": row.line, "part_no": row.part_no, "description": row.description,
            "initial_stock": row.initial_stock, "ct_seconds": row.ct_seconds, "trp": row.trp, "mpq": row.mpq,
            "total_demand": row.total_demand, "avg_daily_demand": round(row.avg_daily_demand, 1),
            "demand": row.demand_by_shift, "plan": row.plan_by_shift, "stock": row.stock_by_shift, "coverage": row.coverage_by_shift,
        }
        for row in matrix.fg_balance
    ]
    sf_balance_data = [
        {
            "site": row.site, "line": row.line, "part_no": row.part_no, "description": row.description,
            "initial_stock": row.initial_stock, "ct_seconds": row.ct_seconds, "trp": row.trp, "mpq": row.mpq,
            "total_demand": row.total_demand, "avg_daily_demand": round(row.avg_daily_demand, 1),
            "demand": row.demand_by_shift, "plan": row.plan_by_shift, "stock": row.stock_by_shift, "coverage": row.coverage_by_shift,
        }
        for row in matrix.sf_balance
    ]

    fg_tasks = [t for t in manual_tasks if t.stage == "FG"]
    sf_tasks = [t for t in manual_tasks if t.stage == "SF"]
    late_tasks = [t for t in fg_tasks if t.end_dt > t.due_datetime]
    on_time_rate = ((len(fg_tasks) - len(late_tasks)) / len(fg_tasks) * 100.0) if fg_tasks else 100.0

    min_start = min((t.start_dt for t in manual_tasks), default=start_dt)
    max_end = max((t.end_dt for t in manual_tasks), default=start_dt + timedelta(days=horizon_days))
    makespan_hours = (max_end - min_start).total_seconds() / 3600.0

    line_utilizations: Dict[str, float] = {}
    for ll in matrix.line_loads:
        if ll.metric == "Loading Rate":
            rates = [v for v in ll.shift_values.values() if isinstance(v, (int, float)) and v > 0]
            line_utilizations[ll.line] = round((sum(rates) / len(rates)) if rates else 0.0, 1)

    total_lines_cnt = len(line_utilizations)
    qualified_lines_cnt = len([r for r in line_utilizations.values() if r >= 95.0])
    avg_line_util = round(sum(line_utilizations.values()) / max(1, total_lines_cnt), 1) if total_lines_cnt else 0.0

    horizon_start = shifts[0].start_dt if shifts else min_start
    horizon_end = shifts[-1].end_dt if shifts else max_end
    total_horizon_sec = max(1.0, (horizon_end - horizon_start).total_seconds())

    task_schedule_list = []
    for t in sorted(manual_tasks, key=lambda x: (x.stage, x.start_dt)):
        slack_h = round((t.due_datetime - t.end_dt).total_seconds() / 3600.0, 1) if t.due_datetime else 999.0
        left_pct = max(0.0, round((t.start_dt - horizon_start).total_seconds() / total_horizon_sec * 100.0, 3))
        dur_pct = max(0.8, round((t.end_dt - t.start_dt).total_seconds() / total_horizon_sec * 100.0, 3))
        due_pct = round((t.due_datetime - horizon_start).total_seconds() / total_horizon_sec * 100.0, 3) if t.due_datetime else None

        task_schedule_list.append({
            "task_id": t.task_id,
            "stage": t.stage,
            "line": t.line,
            "part_no": t.part_no,
            "description": t.description,
            "quantity": int(t.quantity),
            "order_id": t.order_id,
            "ct_seconds": t.ct_seconds,
            "trp": t.trp,
            "mpq": t.mpq,
            "lot_size": t.lot_size,
            "lots": t.lots,
            "start_dt": t.start_dt.strftime("%m/%d %H:%M"),
            "end_dt": t.end_dt.strftime("%m/%d %H:%M"),
            "due_dt": t.due_datetime.strftime("%m/%d %H:%M") if t.due_datetime else "-",
            "duration_hours": round(t.duration_hours, 1),
            "crew_required": getattr(t, "crew_required", 0.0),
            "slack_hours": slack_h,
            "is_late": (t.end_dt > t.due_datetime) if t.due_datetime else False,
            "left_percent": left_pct,
            "width_percent": dur_pct,
            "due_percent": due_pct,
            "shift_breakdown": {f"{k}": v for k, v in t.shift_breakdown.items()},
            "bom_children": [],
            "bom_parents": [],
        })

    return {
        "status": "success",
        "is_manual": True,
        "manual_filename": filename,
        "audit_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "kpi": {
            "status_name": "MANUAL_AUDITED",
            "total_fg_tasks": len(fg_tasks),
            "total_sf_tasks": len(sf_tasks),
            "total_tasks": len(manual_tasks),
            "on_time_delivery_rate": round(on_time_rate, 1),
            "delayed_tasks_count": len(late_tasks),
            "makespan_hours": round(makespan_hours, 1),
            "average_line_utilization": avg_line_util,
            "qualified_lines_count": qualified_lines_cnt,
            "total_lines_count": total_lines_cnt,
            "objective_value": 0,
            "solve_time_seconds": 0.02,
        },
        "shifts": [
            {
                "index": s.index,
                "date_str": s.date_str,
                "shift_name": s.shift_name,
                "is_work_shift": s.is_work_shift,
                "left_percent": max(0.0, round((s.start_dt - horizon_start).total_seconds() / total_horizon_sec * 100.0, 3)),
                "width_percent": max(0.1, round((s.end_dt - s.start_dt).total_seconds() / total_horizon_sec * 100.0, 3)),
            }
            for s in shifts
        ],
        "shift_columns": matrix.shift_columns,
        "line_loads": line_loads_data,
        "fg_balance": fg_balance_data,
        "sf_balance": sf_balance_data,
        "tasks": task_schedule_list,
        "checks": checks_data,
        "compliance_matrix": compliance_rows,
        "empty_shift_reasons": empty_shift_reasons,
        "independent_audit": {
            "passed": independent_audit.passed,
            "summary": f"{len(independent_audit.checks) - len(independent_audit.failed_names())}/{len(independent_audit.checks)} 通过",
            "checks": [
                {
                    "name": check.name,
                    "passed": check.passed,
                    "violations": check.violations[:20],
                    "metrics": check.metrics,
                    "metric_labels": labels_for(check.metrics),
                }
                for check in independent_audit.checks
            ],
        },
        "empty_shift_reasons": empty_shift_reasons,
    }
