"""
Data loader for the APS MVP.

Reads master data, BOM, opening stock and customer demand from the simulation
workbook, then nets demand against stock to derive production tasks: finished
goods rounded up to whole MPQ batches, semi-finished only floored at the MPQ
start quantity.
"""
import csv
import json
import math
import os
import sys
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from src.shift_manager import get_topology


@dataclass
class ProductParam:
    site: str
    line: str
    part_no: str
    description: str
    ct_seconds: float
    trp: float
    mpq: float
    output_per_shift: float
    hours_per_shift: float = 8.0
    shifts_per_day: int = 3
    target_stock: float = 0.0
    target_stock_days: float = 0.0
    lead_time_days: float = 0.0
    scrap_rate: float = 0.0
    headcount: float = 0.0


@dataclass
class CustomerOrder:
    order_id: str
    part_no: str
    description: str
    due_date: str
    due_datetime: datetime
    demand_quantity: float


@dataclass
class ToolingMold:
    site: str
    mold_id: str
    mold_name: str
    applicable_lines: List[str]
    applicable_parts: List[str]
    available_qty: int = 1
    changeover_time_min: float = 30.0
    max_life_strokes: int = 500000
    current_strokes: int = 0
    maintenance_cycle_strokes: int = 50000
    status: str = "可用 (Available)"


@dataclass
class PreventiveMaintenance:
    site: str
    pm_plan_no: str
    line: str
    equipment_name: str
    pm_type: str
    start_date: str
    start_shift: str
    end_date: str
    end_shift: str
    downtime_hours: float = 8.0
    capacity_rate: float = 0.0
    handling_strategy: str = "避让顺延 (Shift_Delay)"
    responsible_team: str = ""
    status: str = "已排定 (Planned)"


@dataclass
class MaterialArrival:
    site: str
    po_number: str
    supplier: str
    part_no: str
    description: str
    eta_date: str
    eta_shift: str
    quantity: float
    qc_release_hours: float = 2.0
    associated_part: str = ""
    status: str = "质检合格 (Released)"


@dataclass
class CustomerTier:
    site: str
    customer_code: str
    customer_name: str
    tier: str  # "STRATEGIC", "TIER_A", "TIER_B"
    priority_weight: int = 10
    otd_target_rate: float = 100.0
    delivery_window: str = "0天严格保供"
    subcontract_allowed: str = "否 (No)"
    associated_products: List[str] = field(default_factory=list)


@dataclass
class ProductionInputSnapshot:
    fg_params: Dict[str, ProductParam]
    sf_params: Dict[str, ProductParam]
    bom: Dict[str, List[Tuple[str, float]]]
    initial_stock: Dict[str, float]
    orders: List[CustomerOrder]
    planning_dates: List[str]
    tooling_molds: List[ToolingMold] = field(default_factory=list)
    pm_schedules: List[PreventiveMaintenance] = field(default_factory=list)
    material_arrivals: List[MaterialArrival] = field(default_factory=list)
    customer_tiers: List[CustomerTier] = field(default_factory=list)


@dataclass
class ProductionTask:
    task_id: str
    stage: str  # "FG" or "SF"
    line: str
    part_no: str
    description: str
    ct_seconds: float
    trp: float
    mpq: int  # FG: rounding multiple; SF: minimum start quantity
    lot_size: int  # units per scheduling lot: MPQ for FG, one piece for SF
    lots: int  # lot_size lots making up quantity
    quantity: float
    lot_millis: int  # one lot's machine time, scaled to whole milliseconds
    order_id: str
    site: str = "WL"
    bom_ratio: float = 1.0
    stock_covered: int = 0
    due_datetime: Optional[datetime] = None
    parent_task_id: Optional[str] = None
    mold_id: Optional[str] = None
    material_arrival_shift: int = 0
    customer_tier: str = "A"
    line_pm_slots: List[int] = field(default_factory=list)
    crew_required: float = 0.0


@dataclass
class StockMovement:
    part_no: str
    stage: str
    opening: float
    produced: float
    consumed: float
    target_stock: float
    closing: float


@dataclass
class RequirementPlan:
    fg_tasks: List[ProductionTask]
    sf_tasks: List[ProductionTask]
    stock_movements: List[StockMovement]
    covered_by_stock: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    invalid_orders: List[Tuple[str, str]] = field(default_factory=list)
    site: str = "WL"
    bom: Dict[str, List[Tuple[str, float]]] = field(default_factory=dict)
    initial_stock: Dict[str, float] = field(default_factory=dict)

    @property
    def tasks(self) -> List[ProductionTask]:
        return self.fg_tasks + self.sf_tasks

    @property
    def has_invalid_orders(self) -> bool:
        return len(self.invalid_orders) > 0

    @property
    def rejected_order_ids(self) -> List[str]:
        return [order_id for order_id, _ in self.invalid_orders]


def _to_float(value, default: float = 0.0) -> float:
    if value is None:
        return default
    s = str(value).strip().replace(',', '')
    if not s:
        return default
    if s.endswith('%'):
        try:
            return float(s[:-1]) / 100.0
        except (TypeError, ValueError):
            return default
    try:
        return float(s)
    except (TypeError, ValueError):
        return default


def snapshot_to_dict(snapshot: ProductionInputSnapshot) -> Dict[str, Any]:
    """Serialize snapshot to pure Python dict suitable for JSON export."""
    d = asdict(snapshot)
    for o in d.get("orders", []):
        if isinstance(o.get("due_datetime"), datetime):
            o["due_datetime"] = o["due_datetime"].isoformat()
    return d


def snapshot_from_dict(d: Dict[str, Any]) -> ProductionInputSnapshot:
    """Deserialize snapshot from Python dict."""
    fg_params = {k: ProductParam(**v) for k, v in d.get("fg_params", {}).items()}
    sf_params = {k: ProductParam(**v) for k, v in d.get("sf_params", {}).items()}
    bom = {k: [(item[0], float(item[1])) for item in v] for k, v in d.get("bom", {}).items()}
    initial_stock = {k: float(v) for k, v in d.get("initial_stock", {}).items()}
    orders = []
    for o in d.get("orders", []):
        dt = o.get("due_datetime")
        if isinstance(dt, str):
            try:
                dt = datetime.fromisoformat(dt)
            except Exception:
                dt = datetime(2026, 10, 20, 15, 0)
        orders.append(CustomerOrder(
            order_id=o["order_id"],
            part_no=o["part_no"],
            description=o["description"],
            due_date=o["due_date"],
            due_datetime=dt,
            demand_quantity=float(o["demand_quantity"]),
        ))
    tooling_molds = [ToolingMold(**m) for m in d.get("tooling_molds", [])]
    pm_schedules = [PreventiveMaintenance(**p) for p in d.get("pm_schedules", [])]
    material_arrivals = [MaterialArrival(**a) for a in d.get("material_arrivals", [])]
    customer_tiers = [CustomerTier(**c) for c in d.get("customer_tiers", [])]
    return ProductionInputSnapshot(
        fg_params=fg_params,
        sf_params=sf_params,
        bom=bom,
        initial_stock=initial_stock,
        orders=orders,
        planning_dates=d.get("planning_dates", []),
        tooling_molds=tooling_molds,
        pm_schedules=pm_schedules,
        material_arrivals=material_arrivals,
        customer_tiers=customer_tiers,
    )


def save_snapshot_json(snapshot: ProductionInputSnapshot, file_path: str) -> None:
    """Save snapshot as readable JSON cache."""
    d = snapshot_to_dict(snapshot)
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=2)


def load_snapshot_json(file_path: str) -> ProductionInputSnapshot:
    """Load snapshot from compiled JSON cache."""
    with open(file_path, "r", encoding="utf-8") as f:
        d = json.load(f)
    return snapshot_from_dict(d)


def load_aps_csv(data_dir: str = "data", site_filter: str = "WL") -> ProductionInputSnapshot:
    """Loads master data, BOM, opening stock and customer demand from CSV files in data_dir,
    generating a self-contained ProductionInputSnapshot and compiling a local snapshot JSON cache.
    """
    if not os.path.isabs(data_dir):
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        candidate = os.path.join(base_dir, data_dir)
        if os.path.exists(candidate):
            data_dir = candidate

    # 1. Load Parameters-FG.csv
    fg_params: Dict[str, ProductParam] = {}
    fg_file = os.path.join(data_dir, "Parameters-FG.csv")
    if os.path.exists(fg_file):
        with open(fg_file, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                if not row or (row.get("Site") or "").strip() != site_filter:
                    continue
                pn = (row.get("Part No.") or row.get("Part No", "")).strip()
                if not pn:
                    continue
                line = (row.get("Production Line") or row.get("Line", "")).strip()
                desc = (row.get("Description") or "").strip()
                ct = _to_float(row.get("Cycle time"))
                trp = _to_float(row.get("Line TRP"), 0.71)
                mpq = _to_float(row.get("Mini Production QTY(MPQ)"), 1.0)
                hours = _to_float(row.get("No. of hours/shift"), 8.0 if site_filter == "WL" else 12.0)
                shifts = int(_to_float(row.get("No. of shift/day"), 3 if site_filter == "WL" else 2))
                output_shift = _to_float(row.get("Output per shift")) or (hours * 3600.0 / ct * trp if ct > 0 else 0.0)
                target_stock = _to_float(row.get("Target Stock"), 0.0)
                fg_params[pn] = ProductParam(
                    site=site_filter,
                    line=line,
                    part_no=pn,
                    description=desc,
                    ct_seconds=ct,
                    trp=trp,
                    mpq=mpq,
                    output_per_shift=output_shift,
                    hours_per_shift=hours,
                    shifts_per_day=shifts,
                    target_stock=target_stock,
                    headcount=_to_float(row.get("Headcount")),
                )

    # 2. Load Parameters-SFN.csv
    sf_params: Dict[str, ProductParam] = {}
    sf_file = os.path.join(data_dir, "Parameters-SFN.csv")
    if os.path.exists(sf_file):
        with open(sf_file, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                if not row or (row.get("Site") or "").strip() != site_filter:
                    continue
                pn = (row.get("Part No.") or row.get("Part No", "")).strip()
                if not pn:
                    continue
                line = (row.get("Production Line") or row.get("Line", "")).strip()
                desc = (row.get("Description") or "").strip()
                ct = _to_float(row.get("Cycle time(S)") or row.get("Cycle time"))
                trp = _to_float(row.get("Line TRP"), 0.85)
                mpq = _to_float(row.get("Mini Production QTY(MPQ)"), 100.0)
                hours = _to_float(row.get("No. of hours/shift"), 8.0 if site_filter == "WL" else 12.0)
                shifts = int(_to_float(row.get("No. of shift/day"), 3 if site_filter == "WL" else 2))
                output_shift = _to_float(row.get("Output per shift")) or (hours * 3600.0 / ct * trp if ct > 0 else 0.0)
                target_stock_days = _to_float(row.get("Target Stock(days)") or row.get("Target Stock (days)"))
                lead_days = _to_float(row.get("生产提前期(days）") or row.get("生产提前期(days)") or row.get("Lead time"))
                scrap = _to_float(row.get("Scrap rate(%)") or row.get("Scrap rate"))
                sf_params[pn] = ProductParam(
                    site=site_filter,
                    line=line,
                    part_no=pn,
                    description=desc,
                    ct_seconds=ct,
                    trp=trp,
                    mpq=mpq,
                    output_per_shift=output_shift,
                    hours_per_shift=hours,
                    shifts_per_day=shifts,
                    target_stock=0.0,
                    target_stock_days=target_stock_days,
                    lead_time_days=lead_days,
                    scrap_rate=scrap,
                    headcount=_to_float(row.get("Head count") or row.get("Headcount")),
                )

    # 3. Load BOM-List.csv
    bom: Dict[str, List[Tuple[str, float]]] = {}
    bom_file = os.path.join(data_dir, "BOM-List.csv")
    if os.path.exists(bom_file):
        with open(bom_file, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                if not row or (row.get("Site") or "").strip() != site_filter:
                    continue
                fg = (row.get("SAP总成") or row.get("成品物料号") or row.get("FG Part No") or "").strip()
                sf = (row.get("半成品号") or row.get("半成品物料号") or row.get("SF Part No") or "").strip()
                ratio = _to_float(row.get("BOM比例") or row.get("Ratio"), 1.0)
                if fg and sf:
                    bom.setdefault(fg, []).append((sf, ratio))

    # 4. Load Initial-Stock.csv
    initial_stock: Dict[str, float] = {}
    stock_file = os.path.join(data_dir, "Initial-Stock.csv")
    if os.path.exists(stock_file):
        with open(stock_file, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                if not row or (row.get("Site") or "").strip() != site_filter:
                    continue
                pn = (row.get("PN") or row.get("Part No") or row.get("物料号") or "").strip()
                stock_val = 0.0
                for k, v in row.items():
                    if "Stock" in k or "库存" in k:
                        stock_val = _to_float(v)
                        break
                if pn:
                    initial_stock[pn] = stock_val

    # 5. Load Customer-Demand.csv
    orders: List[CustomerOrder] = []
    demand_dates: List[str] = []
    dem_file = os.path.join(data_dir, "Customer-Demand.csv")
    if os.path.exists(dem_file):
        with open(dem_file, encoding="utf-8-sig") as f:
            raw_reader = csv.reader(f)
            header_row = None
            date_indices = []
            for row in raw_reader:
                if not row or not any(row):
                    continue
                if header_row is None and any("/" in c for c in row):
                    header_row = [c.strip() for c in row]
                    for idx, c in enumerate(header_row):
                        if "/" in c and any(ch.isdigit() for ch in c):
                            date_indices.append((idx, c))
                            demand_dates.append(c)
                    continue
                if header_row is None:
                    continue
                row_site = row[0].strip() if len(row) > 0 else ""
                if row_site != site_filter:
                    continue
                pn = row[1].strip() if len(row) > 1 else ""
                if not pn:
                    continue
                desc = row[2].strip() if len(row) > 2 else ""
                due_hour = 15 if site_filter == "WL" else 20
                for idx, date_str in date_indices:
                    if idx < len(row):
                        qty = _to_float(row[idx])
                        if qty > 0:
                            try:
                                month, day = (int(p) for p in date_str.split("/"))
                                dt = datetime(2026, month, day, due_hour, 0)
                                o_id = f"ORD_{site_filter}_{pn}_{month:02d}{day:02d}"
                                orders.append(CustomerOrder(
                                    order_id=o_id,
                                    part_no=pn,
                                    description=desc,
                                    due_date=date_str,
                                    due_datetime=dt,
                                    demand_quantity=qty,
                                ))
                            except Exception:
                                pass

    # 6. Load Tooling-Molds.csv
    tooling_molds: List[ToolingMold] = []
    tm_file = os.path.join(data_dir, "Tooling-Molds.csv")
    if os.path.exists(tm_file):
        with open(tm_file, encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                if not row or row[0].strip() in ("Site", "基地 (Site)"):
                    continue
                if row[0].strip() and row[0].strip() != site_filter:
                    continue
                m_id = row[1].strip() if len(row) > 1 else ""
                if not m_id:
                    continue
                lines = [l.strip() for l in (row[3] if len(row) > 3 else "").split(",") if l.strip()]
                parts = [p.strip() for p in (row[4] if len(row) > 4 else "").split(",") if p.strip()]
                tooling_molds.append(ToolingMold(
                    site=row[0].strip() or site_filter,
                    mold_id=m_id,
                    mold_name=row[2].strip() if len(row) > 2 else "",
                    applicable_lines=lines,
                    applicable_parts=parts,
                    available_qty=int(_to_float(row[5] if len(row) > 5 else 1, 1)),
                    changeover_time_min=_to_float(row[6] if len(row) > 6 else 30.0, 30.0),
                    max_life_strokes=int(_to_float(row[7] if len(row) > 7 else 500000, 500000)),
                    current_strokes=int(_to_float(row[8] if len(row) > 8 else 0, 0)),
                    maintenance_cycle_strokes=int(_to_float(row[9] if len(row) > 9 else 50000, 50000)),
                    status=(row[10] if len(row) > 10 else "可用 (Available)").strip(),
                ))

    # 7. Load PM-Schedules.csv
    pm_schedules: List[PreventiveMaintenance] = []
    pm_file = os.path.join(data_dir, "PM-Schedules.csv")
    if os.path.exists(pm_file):
        with open(pm_file, encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                if not row or row[0].strip() in ("Site", "基地 (Site)"):
                    continue
                if row[0].strip() and row[0].strip() != site_filter:
                    continue
                plan_no = row[1].strip() if len(row) > 1 else ""
                if not plan_no:
                    continue
                pm_schedules.append(PreventiveMaintenance(
                    site=row[0].strip() or site_filter,
                    pm_plan_no=plan_no,
                    line=(row[2] if len(row) > 2 else "").strip(),
                    equipment_name=(row[3] if len(row) > 3 else "").strip(),
                    pm_type=(row[4] if len(row) > 4 else "").strip(),
                    start_date=(row[5] if len(row) > 5 else "").strip()[:10],
                    start_shift=(row[6] if len(row) > 6 else "").strip(),
                    end_date=(row[7] if len(row) > 7 else "").strip()[:10],
                    end_shift=(row[8] if len(row) > 8 else "").strip(),
                    downtime_hours=_to_float(row[9] if len(row) > 9 else 8.0, 8.0),
                    capacity_rate=_to_float(row[10] if len(row) > 10 else 0.0, 0.0),
                    handling_strategy=(row[11] if len(row) > 11 else "避让顺延 (Shift_Delay)").strip(),
                    responsible_team=(row[12] if len(row) > 12 else "").strip(),
                    status=(row[13] if len(row) > 13 else "已排定 (Planned)").strip(),
                ))

    # 8. Load Material-Arrivals.csv
    material_arrivals: List[MaterialArrival] = []
    ma_file = os.path.join(data_dir, "Material-Arrivals.csv")
    if os.path.exists(ma_file):
        with open(ma_file, encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                if not row or row[0].strip() in ("Site", "基地 (Site)"):
                    continue
                if row[0].strip() and row[0].strip() != site_filter:
                    continue
                po_no = row[1].strip() if len(row) > 1 else ""
                if not po_no:
                    continue
                material_arrivals.append(MaterialArrival(
                    site=row[0].strip() or site_filter,
                    po_number=po_no,
                    supplier=(row[2] if len(row) > 2 else "").strip(),
                    part_no=(row[3] if len(row) > 3 else "").strip(),
                    description=(row[4] if len(row) > 4 else "").strip(),
                    eta_date=(row[5] if len(row) > 5 else "").strip()[:10],
                    eta_shift=(row[6] if len(row) > 6 else "").strip(),
                    quantity=_to_float(row[7] if len(row) > 7 else 0.0, 0.0),
                    qc_release_hours=_to_float(row[8] if len(row) > 8 else 2.0, 2.0),
                    associated_part=(row[9] if len(row) > 9 else "").strip(),
                    status=(row[10] if len(row) > 10 else "质检合格 (Released)").strip(),
                ))

    # 9. Load Customer-Tiers.csv
    customer_tiers: List[CustomerTier] = []
    ct_file = os.path.join(data_dir, "Customer-Tiers.csv")
    if os.path.exists(ct_file):
        with open(ct_file, encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                if not row or row[0].strip() in ("Site", "基地 (Site)"):
                    continue
                if row[0].strip() and row[0].strip() != site_filter:
                    continue
                c_code = row[1].strip() if len(row) > 1 else ""
                if not c_code:
                    continue
                prods = [p.strip() for p in (row[8] if len(row) > 8 else "").split(",") if p.strip()]
                customer_tiers.append(CustomerTier(
                    site=row[0].strip() or site_filter,
                    customer_code=c_code,
                    customer_name=(row[2] if len(row) > 2 else "").strip(),
                    tier=(row[3] if len(row) > 3 else "TIER_A").strip(),
                    priority_weight=int(_to_float(row[4] if len(row) > 4 else 3, 3)),
                    otd_target_rate=_to_float(row[5] if len(row) > 5 else 98.0, 98.0),
                    delivery_window=(row[6] if len(row) > 6 else "0天严格保供").strip(),
                    subcontract_allowed=(row[7] if len(row) > 7 else "协商外协 (Negotiable)").strip(),
                    associated_products=prods,
                ))

    snapshot = ProductionInputSnapshot(
        fg_params=fg_params,
        sf_params=sf_params,
        bom=bom,
        initial_stock=initial_stock,
        orders=orders,
        planning_dates=demand_dates,
        tooling_molds=tooling_molds,
        pm_schedules=pm_schedules,
        material_arrivals=material_arrivals,
        customer_tiers=customer_tiers,
    )

    try:
        json_path = os.path.join(data_dir, f"snapshot_{site_filter}.json")
        save_snapshot_json(snapshot, json_path)
    except Exception:
        pass

    return snapshot


def _load_aps_excel(excel_path: str, site_filter: str = "WL") -> ProductionInputSnapshot:
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    ct_map: Dict[str, float] = {}
    if "WL CT " in wb.sheetnames:
        ws_ct = wb["WL CT "]
        for r in range(2, ws_ct.max_row + 1):
            material = str(ws_ct.cell(r, 1).value or "").strip()
            ct_value = ws_ct.cell(r, 3).value
            if material and ct_value is not None:
                try:
                    ct_map[material] = float(ct_value)
                except ValueError:
                    pass

    fg_params = _load_params(wb["Parameters-FG"], site_filter, ct_map, default_ct=40.0, default_trp=0.71, default_mpq=1.0)
    sf_params = _load_params(wb["Parameters-SFN"], site_filter, ct_map, default_ct=25.0, default_trp=0.85, default_mpq=100.0, semi_finished=True)

    ws_bom = wb["BOM List"]
    bom: Dict[str, List[Tuple[str, float]]] = {}
    for r in range(2, ws_bom.max_row + 1):
        if str(ws_bom.cell(r, 1).value or "").strip() != site_filter:
            continue
        fg = str(ws_bom.cell(r, 2).value or "").strip()
        sf = str(ws_bom.cell(r, 4).value or "").strip()
        if fg and sf:
            bom.setdefault(fg, []).append((sf, _to_float(ws_bom.cell(r, 6).value, 1.0)))

    ws_stock = wb["Initial Stock"]
    initial_stock: Dict[str, float] = {}
    for r in range(2, ws_stock.max_row + 1):
        if str(ws_stock.cell(r, 1).value or "").strip() != site_filter:
            continue
        pn = str(ws_stock.cell(r, 3).value or "").strip()
        if pn:
            initial_stock[pn] = _to_float(ws_stock.cell(r, 5).value)

    ws_dem = wb["Customer Demand"]
    demand_dates = []
    for c in range(4, ws_dem.max_column + 1):
        v = ws_dem.cell(2, c).value
        if v is None:
            continue
        if isinstance(v, (int, float)) and v > 40000:
            dt_h = datetime(1899, 12, 30) + timedelta(days=int(v))
            demand_dates.append(f"{dt_h.month}/{dt_h.day}")
        elif isinstance(v, datetime):
            demand_dates.append(f"{v.month}/{v.day}")
        else:
            demand_dates.append(str(v).strip())
    orders: List[CustomerOrder] = []
    for r in range(3, ws_dem.max_row + 1):
        if str(ws_dem.cell(r, 1).value or "").strip() != site_filter:
            continue
        pn = str(ws_dem.cell(r, 2).value or "").strip()
        if not pn:
            continue
        desc = str(ws_dem.cell(r, 3).value or "").strip()
        for i, date_str in enumerate(demand_dates):
            qty = _to_float(ws_dem.cell(r, 4 + i).value)
            if qty <= 0:
                continue
            month, day = (int(p) for p in date_str.split("/"))
            orders.append(CustomerOrder(
                order_id=f"ORD_{site_filter}_{pn}_{month:02d}{day:02d}",
                part_no=pn,
                description=desc,
                due_date=date_str,
                due_datetime=datetime(2026, month, day, 15, 0),
                demand_quantity=qty,
            ))

    # Load Tooling-Mold
    tooling_molds: List[ToolingMold] = []
    if "Tooling-Mold" in wb.sheetnames:
        ws_tm = wb["Tooling-Mold"]
        for r in range(2, ws_tm.max_row + 1):
            site_val = str(ws_tm.cell(r, 1).value or "").strip()
            if site_val and site_val != site_filter:
                continue
            m_id = str(ws_tm.cell(r, 2).value or "").strip()
            if not m_id:
                continue
            lines = [l.strip() for l in str(ws_tm.cell(r, 4).value or "").split(",") if l.strip()]
            parts = [p.strip() for p in str(ws_tm.cell(r, 5).value or "").split(",") if p.strip()]
            tooling_molds.append(ToolingMold(
                site=site_val or site_filter,
                mold_id=m_id,
                mold_name=str(ws_tm.cell(r, 3).value or "").strip(),
                applicable_lines=lines,
                applicable_parts=parts,
                available_qty=int(_to_float(ws_tm.cell(r, 6).value, 1)),
                changeover_time_min=_to_float(ws_tm.cell(r, 7).value, 30.0),
                max_life_strokes=int(_to_float(ws_tm.cell(r, 8).value, 500000)),
                current_strokes=int(_to_float(ws_tm.cell(r, 9).value, 0)),
                maintenance_cycle_strokes=int(_to_float(ws_tm.cell(r, 10).value, 50000)),
                status=str(ws_tm.cell(r, 11).value or "可用 (Available)").strip(),
            ))

    # Load Preventive-Maintenance
    pm_schedules: List[PreventiveMaintenance] = []
    if "Preventive-Maintenance" in wb.sheetnames:
        ws_pm = wb["Preventive-Maintenance"]
        for r in range(2, ws_pm.max_row + 1):
            site_val = str(ws_pm.cell(r, 1).value or "").strip()
            if site_val and site_val != site_filter:
                continue
            plan_no = str(ws_pm.cell(r, 2).value or "").strip()
            if not plan_no:
                continue
            pm_schedules.append(PreventiveMaintenance(
                site=site_val or site_filter,
                pm_plan_no=plan_no,
                line=str(ws_pm.cell(r, 3).value or "").strip(),
                equipment_name=str(ws_pm.cell(r, 4).value or "").strip(),
                pm_type=str(ws_pm.cell(r, 5).value or "").strip(),
                start_date=str(ws_pm.cell(r, 6).value or "").strip()[:10],
                start_shift=str(ws_pm.cell(r, 7).value or "").strip(),
                end_date=str(ws_pm.cell(r, 8).value or "").strip()[:10],
                end_shift=str(ws_pm.cell(r, 9).value or "").strip(),
                downtime_hours=_to_float(ws_pm.cell(r, 10).value, 8.0),
                capacity_rate=_to_float(ws_pm.cell(r, 11).value, 0.0),
                handling_strategy=str(ws_pm.cell(r, 12).value or "避让顺延 (Shift_Delay)").strip(),
                responsible_team=str(ws_pm.cell(r, 13).value or "").strip(),
                status=str(ws_pm.cell(r, 14).value or "已排定 (Planned)").strip(),
            ))

    # Load Material-Arrival-ETA
    material_arrivals: List[MaterialArrival] = []
    if "Material-Arrival-ETA" in wb.sheetnames:
        ws_arr = wb["Material-Arrival-ETA"]
        for r in range(2, ws_arr.max_row + 1):
            site_val = str(ws_arr.cell(r, 1).value or "").strip()
            if site_val and site_val != site_filter:
                continue
            po_no = str(ws_arr.cell(r, 2).value or "").strip()
            if not po_no:
                continue
            material_arrivals.append(MaterialArrival(
                site=site_val or site_filter,
                po_number=po_no,
                supplier=str(ws_arr.cell(r, 3).value or "").strip(),
                part_no=str(ws_arr.cell(r, 4).value or "").strip(),
                description=str(ws_arr.cell(r, 5).value or "").strip(),
                eta_date=str(ws_arr.cell(r, 6).value or "").strip()[:10],
                eta_shift=str(ws_arr.cell(r, 7).value or "").strip(),
                quantity=_to_float(ws_arr.cell(r, 8).value, 0.0),
                qc_release_hours=_to_float(ws_arr.cell(r, 9).value, 2.0),
                associated_part=str(ws_arr.cell(r, 10).value or "").strip(),
                status=str(ws_arr.cell(r, 11).value or "质检合格 (Released)").strip(),
            ))

    # Load Customer-Priority
    customer_tiers: List[CustomerTier] = []
    if "Customer-Priority" in wb.sheetnames:
        ws_cust = wb["Customer-Priority"]
        for r in range(2, ws_cust.max_row + 1):
            site_val = str(ws_cust.cell(r, 1).value or "").strip()
            if site_val and site_val != site_filter:
                continue
            c_code = str(ws_cust.cell(r, 2).value or "").strip()
            if not c_code:
                continue
            prods = [p.strip() for p in str(ws_cust.cell(r, 9).value or "").split(",") if p.strip()]
            customer_tiers.append(CustomerTier(
                site=site_val or site_filter,
                customer_code=c_code,
                customer_name=str(ws_cust.cell(r, 3).value or "").strip(),
                tier=str(ws_cust.cell(r, 4).value or "TIER_A").strip(),
                priority_weight=int(_to_float(ws_cust.cell(r, 5).value, 3)),
                otd_target_rate=_to_float(ws_cust.cell(r, 6).value, 98.0),
                delivery_window=str(ws_cust.cell(r, 7).value or "0天严格保供").strip(),
                subcontract_allowed=str(ws_cust.cell(r, 8).value or "协商外协 (Negotiable)").strip(),
                associated_products=prods,
            ))

    return ProductionInputSnapshot(
        fg_params=fg_params,
        sf_params=sf_params,
        bom=bom,
        initial_stock=initial_stock,
        orders=orders,
        planning_dates=demand_dates,
        tooling_molds=tooling_molds,
        pm_schedules=pm_schedules,
        material_arrivals=material_arrivals,
        customer_tiers=customer_tiers,
    )


def load_aps_database(source_path: str = "data", site_filter: str = "WL") -> ProductionInputSnapshot:
    """Unified entry point to load APS dataset.
    Prioritizes loading from explicit .xlsx file if given, otherwise loads from CSV directory.
    """
    # If source_path is an existing xlsx file
    if os.path.isfile(source_path) and source_path.lower().endswith(".xlsx"):
        return _load_aps_excel(source_path, site_filter=site_filter)

    if os.path.isdir(source_path):
        return load_aps_csv(source_path, site_filter=site_filter)

    # Check candidate data dir
    candidate_data = os.path.join(os.path.dirname(os.path.abspath(source_path)), "data")
    if os.path.exists(candidate_data) and os.path.isdir(candidate_data):
        return load_aps_csv(candidate_data, site_filter=site_filter)

    # Default fallback to 'data'
    if os.path.isdir("data"):
        return load_aps_csv("data", site_filter=site_filter)

    raise FileNotFoundError(f"Cannot load APS dataset: invalid path '{source_path}'")


def _load_params(ws, site_filter: str, ct_map: Dict[str, float], *, default_ct: float,
                 default_trp: float, default_mpq: float, semi_finished: bool = False) -> Dict[str, ProductParam]:
    params: Dict[str, ProductParam] = {}
    hours_col = 11 if semi_finished else 9
    shifts_col = 12 if semi_finished else 10
    headcount_col = 14 if semi_finished else 12
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() != site_filter:
            continue
        pn = str(ws.cell(r, 4).value or "").strip()
        if not pn:
            continue
        ct = _to_float(ws.cell(r, 7).value) or ct_map.get(pn, default_ct)
        trp = _to_float(ws.cell(r, 3).value, default_trp) or default_trp
        mpq = _to_float(ws.cell(r, 6).value, default_mpq) or default_mpq
        hours = _to_float(ws.cell(r, hours_col).value, 8.0) or 8.0
        params[pn] = ProductParam(
            site=site_filter,
            line=str(ws.cell(r, 2).value or "").strip(),
            part_no=pn,
            description=str(ws.cell(r, 5).value or "").strip(),
            ct_seconds=ct,
            trp=trp,
            mpq=mpq,
            output_per_shift=hours * 3600.0 / ct * trp,
            hours_per_shift=hours,
            shifts_per_day=int(_to_float(ws.cell(r, shifts_col).value, 3.0) or 3),
            target_stock=0.0 if semi_finished else _to_float(ws.cell(r, 8).value),
            target_stock_days=_to_float(ws.cell(r, 8).value) if semi_finished else 0.0,
            lead_time_days=_to_float(ws.cell(r, 9).value) if semi_finished else 0.0,
            scrap_rate=_to_float(ws.cell(r, 10).value) if semi_finished else 0.0,
            headcount=_to_float(ws.cell(r, headcount_col).value),
        )
    return params


TIME_SCALE = 1000


def lot_millis(lot_size: float, ct_seconds: float, trp: float) -> int:
    """Milliseconds for one production lot: CP-SAT only takes integer time."""
    return max(1, int(round(lot_size * ct_seconds / trp * TIME_SCALE)))


def build_requirement_plan(snapshot: ProductionInputSnapshot) -> RequirementPlan:
    """Net customer demand against opening and safety stock, then explode BOM.

    Finished goods are produced in whole MPQ batches, so the rounding surplus stays
    on hand and offsets the next order for the same part. Semi-finished output only
    has to clear the MPQ start quantity, so an order of 120 pieces costs 120 pieces
    instead of two batches of 100.
    """
    site = next((p.site for p in snapshot.fg_params.values()), "WL")
    topology = get_topology(site)

    fg_tasks: List[ProductionTask] = []
    sf_tasks: List[ProductionTask] = []
    covered_by_stock: List[str] = []
    warnings: List[str] = []
    invalid_orders: List[Tuple[str, str]] = []
    crew_warned: set = set()

    fg_avail = {pn: snapshot.initial_stock.get(pn, 0.0) for pn in snapshot.fg_params}
    sf_avail = {pn: snapshot.initial_stock.get(pn, 0.0) for pn in snapshot.sf_params}
    fg_produced: Dict[str, float] = {}
    sf_produced: Dict[str, float] = {}
    fg_consumed: Dict[str, float] = {}
    sf_consumed: Dict[str, float] = {}

    for pn, param in snapshot.sf_params.items():
        if param.target_stock_days > 0:
            warnings.append(
                f"{pn}: 半成品安全库存以天数({param.target_stock_days}天)给出，尚未折算为件数，本次按 0 处理"
            )

    # High-order master data lookup maps
    mold_map: Dict[str, str] = {}
    for tm in snapshot.tooling_molds:
        for p in tm.applicable_parts:
            mold_map[p] = tm.mold_id

    cust_tier_map: Dict[str, str] = {}
    for ct in snapshot.customer_tiers:
        tier_code = "STRATEGIC" if "STRATEGIC" in ct.tier.upper() else ("TIER_B" if "TIER_B" in ct.tier.upper() or "B级" in ct.tier else "TIER_A")
        for p in ct.associated_products:
            cust_tier_map[p] = tier_code

    mat_arrival_map: Dict[str, int] = {}
    for ma in snapshot.material_arrivals:
        if ma.eta_date:
            try:
                parsed_date = datetime.strptime(ma.eta_date.replace("-", "/"), "%Y/%m/%d") if len(ma.eta_date) > 5 else datetime.strptime(ma.eta_date.replace("-", "/"), "%m/%d")
                d_idx = -1
                for idx, pdate in enumerate(snapshot.planning_dates):
                    p_parsed = datetime.strptime(pdate, "%m/%d")
                    if p_parsed.month == parsed_date.month and p_parsed.day == parsed_date.day:
                        d_idx = idx
                        break
                if d_idx >= 0:
                    target_shift = topology.slot_index(d_idx, ma.eta_shift)
                    part_target = ma.associated_part or ma.part_no
                    mat_arrival_map[part_target] = max(mat_arrival_map.get(part_target, 0), target_shift)
            except Exception:
                pass

    line_pm_slots: Dict[str, List[int]] = defaultdict(list)
    for pm in snapshot.pm_schedules:
        if pm.capacity_rate == 0 or "避让" in pm.handling_strategy or "切分" in pm.handling_strategy:
            try:
                start_parsed = datetime.strptime(pm.start_date.replace("-", "/"), "%Y/%m/%d") if len(pm.start_date) > 5 else datetime.strptime(pm.start_date.replace("-", "/"), "%m/%d")
                end_parsed = datetime.strptime(pm.end_date.replace("-", "/"), "%Y/%m/%d") if len(pm.end_date) > 5 else datetime.strptime(pm.end_date.replace("-", "/"), "%m/%d")

                d_start = 0
                for idx, pdate in enumerate(snapshot.planning_dates):
                    p_parsed = datetime.strptime(pdate, "%m/%d")
                    if p_parsed.month == start_parsed.month and p_parsed.day == start_parsed.day:
                        d_start = idx
                        break
                d_end = d_start
                for idx, pdate in enumerate(snapshot.planning_dates):
                    p_parsed = datetime.strptime(pdate, "%m/%d")
                    if p_parsed.month == end_parsed.month and p_parsed.day == end_parsed.day:
                        d_end = idx
                        break

                s_start = topology.slot_index(d_start, pm.start_shift)
                s_end = topology.slot_index(d_end, pm.end_shift)
                for s in range(s_start, s_end + 1):
                    if s not in line_pm_slots[pm.line]:
                        line_pm_slots[pm.line].append(s)
            except Exception:
                pass

    for order in sorted(snapshot.orders, key=lambda o: (o.due_datetime, o.part_no)):
        param = snapshot.fg_params.get(order.part_no)
        if param is None:
            warnings.append(f"{order.order_id}: 成品 {order.part_no} 无 Parameters-FG 主数据，跳过")
            invalid_orders.append((order.order_id, f"成品 {order.part_no} 缺失主数据"))
            continue
        if param.headcount <= 0 and order.part_no not in crew_warned:
            warnings.append(f"{order.part_no}: 缺少 Headcount 主数据，人力工时按 0 计")
            crew_warned.add(order.part_no)

        net_need = max(0.0, order.demand_quantity + param.target_stock - fg_avail[order.part_no])
        fg_consumed[order.part_no] = fg_consumed.get(order.part_no, 0.0) + order.demand_quantity

        if net_need <= 0:
            fg_avail[order.part_no] -= order.demand_quantity
            covered_by_stock.append(
                f"{order.order_id}: 需求 {order.demand_quantity:.0f} 全部由库存覆盖"
                f"（期初 {snapshot.initial_stock.get(order.part_no, 0.0):.0f}），不排产"
            )
            continue

        batches = int(math.ceil(net_need / param.mpq))
        lot_size = int(round(param.mpq))
        qty = batches * lot_size
        fg_avail[order.part_no] += qty - order.demand_quantity
        fg_produced[order.part_no] = fg_produced.get(order.part_no, 0.0) + qty

        fg_task = ProductionTask(
            task_id=f"FG::{order.order_id}",
            stage="FG",
            line=param.line,
            part_no=order.part_no,
            description=param.description,
            ct_seconds=param.ct_seconds,
            trp=param.trp,
            mpq=lot_size,
            lot_size=lot_size,
            lots=batches,
            quantity=float(qty),
            lot_millis=lot_millis(lot_size, param.ct_seconds, param.trp),
            order_id=order.order_id,
            site=site,
            due_datetime=order.due_datetime,
            mold_id=mold_map.get(order.part_no),
            material_arrival_shift=mat_arrival_map.get(order.part_no, 0),
            customer_tier=cust_tier_map.get(order.part_no, "TIER_A"),
            line_pm_slots=line_pm_slots.get(param.line, []),
            crew_required=param.headcount,
        )
        fg_tasks.append(fg_task)

        for sf_pn, ratio in snapshot.bom.get(order.part_no, []):
            sf_param = snapshot.sf_params.get(sf_pn)
            if sf_param is None:
                warnings.append(f"{fg_task.task_id}: BOM 半成品 {sf_pn} 无 Parameters-SFN 主数据，跳过齐套")
                invalid_orders.append((order.order_id, f"半成品 {sf_pn} 缺失主数据"))
                continue
            if sf_param.headcount <= 0 and sf_pn not in crew_warned:
                warnings.append(f"{sf_pn}: 缺少 Headcount 主数据，人力工时按 0 计")
                crew_warned.add(sf_pn)

            gross_need = qty * ratio
            sf_usable = max(0.0, sf_avail[sf_pn] - sf_param.target_stock)
            sf_net = max(0.0, gross_need + sf_param.target_stock - sf_avail[sf_pn])
            sf_consumed[sf_pn] = sf_consumed.get(sf_pn, 0.0) + gross_need

            if sf_net <= 0:
                sf_avail[sf_pn] -= gross_need
                continue

            min_lot = int(round(sf_param.mpq))
            sf_qty = max(min_lot, int(math.ceil(sf_net)))
            sf_avail[sf_pn] += sf_qty - gross_need
            sf_produced[sf_pn] = sf_produced.get(sf_pn, 0.0) + sf_qty

            sf_tasks.append(ProductionTask(
                task_id=f"SF::{order.order_id}::{sf_pn}",
                stage="SF",
                line=sf_param.line,
                part_no=sf_pn,
                description=sf_param.description,
                ct_seconds=sf_param.ct_seconds,
                trp=sf_param.trp,
                mpq=min_lot,
                lot_size=1,
                lots=sf_qty,
                quantity=float(sf_qty),
                lot_millis=lot_millis(1, sf_param.ct_seconds, sf_param.trp),
                order_id=order.order_id,
                site=site,
                bom_ratio=ratio,
                stock_covered=int(round(sf_usable)),
                due_datetime=order.due_datetime,
                parent_task_id=fg_task.task_id,
                mold_id=mold_map.get(sf_pn),
                material_arrival_shift=mat_arrival_map.get(sf_pn, 0),
                customer_tier=cust_tier_map.get(order.part_no, "TIER_A"),
                line_pm_slots=line_pm_slots.get(sf_param.line, []),
                crew_required=sf_param.headcount,
            ))

    movements = [
        StockMovement(
            part_no=pn,
            stage="FG",
            opening=snapshot.initial_stock.get(pn, 0.0),
            produced=fg_produced.get(pn, 0.0),
            consumed=fg_consumed.get(pn, 0.0),
            target_stock=snapshot.fg_params[pn].target_stock,
            closing=fg_avail[pn],
        )
        for pn in snapshot.fg_params
        if fg_produced.get(pn) or fg_consumed.get(pn) or snapshot.initial_stock.get(pn)
    ] + [
        StockMovement(
            part_no=pn,
            stage="SF",
            opening=snapshot.initial_stock.get(pn, 0.0),
            produced=sf_produced.get(pn, 0.0),
            consumed=sf_consumed.get(pn, 0.0),
            target_stock=snapshot.sf_params[pn].target_stock,
            closing=sf_avail[pn],
        )
        for pn in snapshot.sf_params
        if sf_produced.get(pn) or sf_consumed.get(pn) or snapshot.initial_stock.get(pn)
    ]

    return RequirementPlan(
        fg_tasks=fg_tasks,
        sf_tasks=sf_tasks,
        stock_movements=movements,
        covered_by_stock=covered_by_stock,
        warnings=warnings,
        invalid_orders=invalid_orders,
        site=site,
        bom=dict(snapshot.bom),
        initial_stock=dict(snapshot.initial_stock),
    )
