"""
Rule registry for the APS MVP (Unified Multi-Site Architecture: WL 27 Rules & SY38 9 Rules).

Source of truth: '排产原则-约束条件' sheet in 'APS Database- Simulation draft.xlsx'.
Classifies rules by site, category, nature, and landing layer.
"""
import os
from dataclasses import dataclass
from enum import Enum
from typing import Dict, List, Optional
import openpyxl


class ConstraintCategory(str, Enum):
    EQUIPMENT_CAPACITY = "设备与产能"
    CORE_SCHEDULING = "核心排班"
    MATERIAL_INVENTORY = "物料与库存"
    DEMAND_BACKWARD = "需求倒排"
    EXCEPTION_RESPONSE = "异常应对"
    SYSTEM_ARCHITECTURE = "系统架构"
    TEST_SCENARIO = "测试场景"
    TOOLING_PERSONNEL = "工装与人员"
    ORDER_DELIVERY = "订单与交期"


class ConstraintNature(str, Enum):
    HARD = "硬约束"
    SOFT = "软约束"
    STRATEGY = "业务策略"
    ARCHITECTURE = "架构约束"
    TEST_CASE = "测试用例"


# Backward compatibility alias
RuleCategory = ConstraintNature


class RuleStatus(str, Enum):
    IMPLEMENTED = "PASS"             # Fully implemented & verified
    PARTIAL = "PARTIAL"             # Partially implemented (contains open assumptions)
    UNIMPLEMENTED = "TODO"          # Not implemented yet
    BLOCKED_DATA = "N.A.(数据缺失)"  # Blocked by missing master data fields
    VIOLATED = "VIOLATED"           # Violated by current code design


@dataclass(frozen=True)
class Rule:
    rule_id: str
    row_index: int
    stage: str
    category: ConstraintCategory
    nature: ConstraintNature
    name: str
    weight: Optional[float]
    status: RuleStatus
    landing_layer: str
    validation_check: Optional[str]
    decision_id: Optional[str]
    description: str = ""
    ai_logic: Optional[str] = None


def load_rules_from_excel(excel_path: Optional[str] = None) -> Dict[str, Dict[str, Rule]]:
    if not excel_path or not os.path.exists(excel_path):
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        excel_path = os.path.join(base_dir, "dataSource.xlsx")
        if not os.path.exists(excel_path):
            excel_path = os.path.join(base_dir, "APS Database- Simulation draft.xlsx")
    if not os.path.exists(excel_path):
        return {"WL": REGISTRY_WE34_FALLBACK, "SY38": REGISTRY_SY38_FALLBACK}

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = [s for s in wb.worksheets if "排产原则" in s.title][0]
        cat_map = {c.value: c for c in ConstraintCategory}
        nature_map = {n.value: n for n in ConstraintNature}
        rules_by_site: Dict[str, Dict[str, Rule]] = {"WL": {}, "SY38": {}}
        we_idx = 1
        sy_idx = 1
        for r in range(2, ws.max_row + 1):
            site = str(ws.cell(r, 1).value or "").strip()
            mat_type = str(ws.cell(r, 2).value or "").strip()
            cat_str = str(ws.cell(r, 3).value or "").strip()
            name = str(ws.cell(r, 4).value or "").strip()
            nature_str = str(ws.cell(r, 5).value or "").strip()
            weight_val = ws.cell(r, 6).value
            logic = str(ws.cell(r, 7).value or "").strip()
            if not name:
                continue
            category = cat_map.get(cat_str, ConstraintCategory.CORE_SCHEDULING)
            nature = nature_map.get(nature_str, ConstraintNature.HARD)
            weight = float(weight_val) if isinstance(weight_val, (int, float)) else None
            if site in ("WE34", "WL"):
                target_site = "WL"
                rule_id = f"R-WL-{we_idx:02d}"
                we_idx += 1
            elif site in ("SY38", "SY"):
                target_site = "SY38"
                rule_id = f"R-SY-{sy_idx:02d}"
                sy_idx += 1
            else:
                continue
            rules_by_site[target_site][rule_id] = Rule(
                rule_id=rule_id,
                row_index=r,
                stage=mat_type or "FG&SFIN",
                category=category,
                nature=nature,
                name=name,
                weight=weight,
                status=RuleStatus.IMPLEMENTED,
                landing_layer="排产引擎核心约束层",
                validation_check=f"规则 {rule_id} 独立审计",
                decision_id=None,
                description=name,
                ai_logic=logic,
            )
        return rules_by_site
    except Exception as e:
        print(f"Warning: Failed to load rules dynamically from Excel: {e}")
        return {"WL": REGISTRY_WE34_FALLBACK, "SY38": REGISTRY_SY38_FALLBACK}


REGISTRY_WE34_FALLBACK: Dict[str, Rule] = {
    "R-WL-01": Rule(
        rule_id="R-WL-01",
        row_index=2,
        stage="FG&SFIN",
        category=ConstraintCategory.TOOLING_PERSONNEL,
        nature=ConstraintNature.HARD,
        name="多设备共用的同一模具/工装具备物理独占性，严格禁止跨设备时段冲突",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-01 独立审计",
        decision_id=None,
        description="多设备共用的同一模具/工装具备物理独占性，严格禁止跨设备时段冲突",
        ai_logic="【逻辑公式】工装模具独占硬约束：若 Mold(Task_i) == Mold(Task_j) 且 i ≠ j，则 [Start_Time(i), End_Time(i)] ∩ [Start_Time(j), End_Time(j)] = ∅。\n【业务边界】共用模具/夹具物理唯一，禁止跨设备并发排产，强制时间轴序列化。\n【场景示例】设备Line 1与Line 2共用模具M01。Line 1安排周二08:00-12:00生产，Line 2若需使用M01，排产最早只能从周二12:00（含换模搬运时间）后开始排入。",
    ),
    "R-WL-02": Rule(
        rule_id="R-WL-02",
        row_index=3,
        stage="FG&SFIN",
        category=ConstraintCategory.EXCEPTION_RESPONSE,
        nature=ConstraintNature.STRATEGY,
        name="面对前紧后松等负荷极端不均场景，启用跨周期滚动与弹性动态日历",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-02 独立审计",
        decision_id=None,
        description="面对前紧后松等负荷极端不均场景，启用跨周期滚动与弹性动态日历",
        ai_logic="【异常策略】非均衡波峰响应：支持宏观日历跨周期滚动切片与局部日历弹性分裂调度。\n【业务边界】应对月初集中爆单、月末断崖式闲置等“前端拥堵”极端场景。\n【仿真示例】某月前两周需求10000件（极限产能仅7000件），后两周需求仅1000件。系统不采用全月均摊做五休二，而是前两周启动无休连排+跨月预排+缺口延期组合策略。",
    ),
    "R-WL-03": Rule(
        rule_id="R-WL-03",
        row_index=4,
        stage="FG&SFIN",
        category=ConstraintCategory.EXCEPTION_RESPONSE,
        nature=ConstraintNature.STRATEGY,
        name="远期产能不足而近期富余时，自动将远期缺口需求前置到近期空闲时段提前生产（削峰填谷）",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-03 独立审计",
        decision_id=None,
        description="远期产能不足而近期富余时，自动将远期缺口需求前置到近期空闲时段提前生产（削峰填谷）",
        ai_logic="【排产策略】跨期预产与产能平准化：若 Demand_Far > Capacity_Far 且 Demand_Near < Capacity_Near，则将远期缺口 Gap_Far 前置到近期空闲窗口提前生产。\n【业务边界】通过削峰填谷平准化产能负荷，防范远期交付挤兑与近期机台闲置。\n【场景示例】某月后半月订单严重超载（缺口2000件），而前半月尚有3000件富余产能。系统自动将后半月2000件工单前置至前半月空闲时段提前排产备库，实现跨期平准化。",
    ),
    "R-WL-04": Rule(
        rule_id="R-WL-04",
        row_index=5,
        stage="FG&SFIN",
        category=ConstraintCategory.EXCEPTION_RESPONSE,
        nature=ConstraintNature.STRATEGY,
        name="支持同月内分段设置周/日级弹性日历，前段峰值连产、后段低谷轮休",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-04 独立审计",
        decision_id=None,
        description="支持同月内分段设置周/日级弹性日历，前段峰值连产、后段低谷轮休",
        ai_logic="【异常策略】分段日历粒度：Calendar(Line, Period_1) ≠ Calendar(Line, Period_2)，支持同月多套日历组合。\n【业务边界】严禁同月一刀切排班，允许前紧后松动态切换。\n【仿真示例】同一产线在4月第1-2周绑定“7天×24小时连续生产日历”；进入第3-4周订单骤减，系统自动切换绑定为“做三休四集中轮休日历”。",
    ),
    "R-WL-05": Rule(
        rule_id="R-WL-05",
        row_index=6,
        stage="FG&SFIN",
        category=ConstraintCategory.EXCEPTION_RESPONSE,
        nature=ConstraintNature.STRATEGY,
        name="极限加班仍无法满足交付时锁定产能缺口，自动输出交期延期平移建议",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-05 独立审计",
        decision_id=None,
        description="极限加班仍无法满足交付时锁定产能缺口，自动输出交期延期平移建议",
        ai_logic="【异常策略】产能锁定与延期建议：若 极限产能 + 预排 + 加班 仍无法满足交期，锁定超额订单并生成《交期调整建议表》。\n【业务边界】不可抗力交付缺口透明化，算法自动给出最佳延期平移建议日期。\n【场景示例】某紧急订单1500件交期为10日，产线全负荷运转至10日仍有500件绝对缺口。系统锁定该500件并出具告警建议：“建议顺延至20日（月末空闲期）交付”。",
    ),
    "R-WL-06": Rule(
        rule_id="R-WL-06",
        row_index=7,
        stage="FG&SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.SOFT,
        name="极度空闲时（周需求≤极限产能4/7），启用做四休三（周一至周四排产）",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-06 独立审计",
        decision_id=None,
        description="极度空闲时（周需求≤极限产能4/7），启用做四休三（周一至周四排产）",
        ai_logic="【排班规则】阶梯1（极度空闲）：若 Demand_Weekly ≤ 4/7 × Max_Capacity，Work_Days = {周一, 周二, 周三, 周四}，Off_Days = {周五, 周六, 周日}。\n【业务边界】低负荷时集中压缩开机天数，降本节能。\n【仿真示例】周极限产能7000件。当周总需求仅3500件（≤4000件），系统自动激活“做四休三”日历，周一至周四满负荷生产，周五至周日停工休假。",
    ),
    "R-WL-07": Rule(
        rule_id="R-WL-07",
        row_index=8,
        stage="FG&SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.SOFT,
        name="正常均衡时（周需求在产能4/7至5/7），启用标准做五休二工作制",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-07 独立审计",
        decision_id=None,
        description="正常均衡时（周需求在产能4/7至5/7），启用标准做五休二工作制",
        ai_logic="【排班规则】阶梯2（正常均衡）：若 4/7 × Max_Cap < Demand_Weekly ≤ 5/7 × Max_Cap，Work_Days = {周一至周五}，Off_Days = {周六, 周日}。\n【业务边界】标准法定工作制，工时平准化排程。\n【仿真示例】周极限产能7000件。周需求为4800件（在4000与5000之间），系统激活“做五休二”日历，周一至周五平摊生产960件/天，周六周日产线正常双休。",
    ),
    "R-WL-08": Rule(
        rule_id="R-WL-08",
        row_index=9,
        stage="FG&SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.SOFT,
        name="产能紧张时（周需求>极限产能5/7），取消周末排休，启动连续生产",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-08 独立审计",
        decision_id=None,
        description="产能紧张时（周需求>极限产能5/7），取消周末排休，启动连续生产",
        ai_logic="【排班规则】阶梯3（产能紧张）：若 Demand_Weekly > 5/7 × Max_Cap，Work_Days = {周一至周日}，Work_Calendar = 连续生产/无休。\n【业务边界】订单高企时打破常规周末休息，全量释放极限产能。\n【仿真示例】周极限产能7000件。周需求达6200件（>5000件），系统触发排班升阶，自动取消周六周日休息，排定7天连续开工以保障准时交付。",
    ),
    "R-WL-09": Rule(
        rule_id="R-WL-09",
        row_index=10,
        stage="FING",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.HARD,
        name="批次量必须满足MPQ，单班不足时强制跨班连续生产严禁切断换线",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-09 独立审计",
        decision_id=None,
        description="批次量必须满足MPQ，单班不足时强制跨班连续生产严禁切断换线",
        ai_logic="【逻辑公式】连续跨班不可分割硬约束：Batch ≥ MPQ；若 班次产能 < MPQ，连续排产不可换线。\n【业务边界】整批生产不可中途中断，单班做不完强制在下一班连续排满MPQ。\n【场景示例】MPQ为1200件，产线单班仅能做800件。系统自动锁定早班800件+中班400件连续作业，中间严禁切换其他产品。",
    ),
    "R-WL-10": Rule(
        rule_id="R-WL-10",
        row_index=11,
        stage="FG&SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.HARD,
        name="成品装配时间围栏内已排产工单自动锁定，系统重排不得擅自挪移",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-10 独立审计",
        decision_id=None,
        description="成品装配时间围栏内已排产工单自动锁定，系统重排不得擅自挪移",
        ai_logic="【逻辑公式】装配围栏保护硬约束：若 Schedule_Time(j) ∈ [Now, Now + Fence_Period]，Auto_Reschedule(j) = 0。\n【业务边界】锁定总装近端已发工单，杜绝系统频繁变动引发线边物料与人员调配混乱。\n【场景示例】成品总装线设定近端12小时时间围栏。已排产上线并配料的工单被强制锁定，系统重排时不得自动插单、挪移或撤销，修改仅限计划员特批手工调整。",
    ),
    "R-WL-11": Rule(
        rule_id="R-WL-11",
        row_index=12,
        stage="FING&SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.SOFT,
        name="保交付前提下同类订单合并换型趋零，且跨班交接优先继承前班模具严防额外换型（权重80%）",
        weight=0.8,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-11 独立审计",
        decision_id=None,
        description="保交付前提下同类订单合并换型趋零，且跨班交接优先继承前班模具严防额外换型（权重80%）",
        ai_logic="【逻辑公式】换型成本与跨班继承软约束（权重80%）：Min ∑ [ 0.80 × Setup_Penalty × 𝕀(Setup_Type(t) ≠ Setup_Type(t-1)) ]；优先保持 Shift_Handover_Setup(t) = Shift_Handover_Setup(t-1)。\n【业务边界】在保交付前提下，极小化单班内及跨班换型损耗，避免交接班因打断产生二次换模。\n【场景示例】班次交接时，前班正在生产产品A（模具M1）。接班次优先接续排产A产品，直至该工单或合并批次完成，严禁无故在接班初始强行切换模具M2。",
    ),
    "R-WL-12": Rule(
        rule_id="R-WL-12",
        row_index=13,
        stage="FING&SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.SOFT,
        name="负荷不足95%时仅顺位前拉紧邻工单且允许跨班连续作业，严禁跳跃拣单（权重80%）",
        weight=0.8,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-12 独立审计",
        decision_id=None,
        description="负荷不足95%时仅顺位前拉紧邻工单且允许跨班连续作业，严禁跳跃拣单（权重80%）",
        ai_logic="【逻辑公式】顺位填补与跨班连续软约束（权重80%）：若 Utilization(Shift_t) < 95%，Pull_Order = Order_{Next}(t) 且 跨越班次时 Batch_Merged = 1；禁止跨序拣单 Forbidden(Order_{t+k}, k > 1)。\n【业务边界】防凑单碎片化与频繁断线，保持时间轴完整性，拉取订单跨班保持作业连续性。\n【场景示例】白班负荷仅80%，时间轴紧邻的中班订单需生产4小时。系统将该订单顺位提前至白班末尾开始生产并平滑顺延至中班，两班合并为单一连续批次，禁止跳单挑拣后天急件。",
    ),
    "R-WL-13": Rule(
        rule_id="R-WL-13",
        row_index=14,
        stage="FING&SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.HARD,
        name="成品生产精准匹配总订单需求，累计达标后即刻封盘停排",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-13 独立审计",
        decision_id=None,
        description="成品生产精准匹配总订单需求，累计达标后即刻封盘停排",
        ai_logic="【逻辑公式】成品按单交付硬约束：∑ Output = Total_Demand，严格闭环生产。\n【业务边界】杜绝过度推式生产造成成品高库龄积压，匹配总订单即止。\n【场景示例】某料号总采购订单2000件。系统精准排满2000件后即刻关停工单，不再排入任何无订单预测的额外数量。",
    ),
    "R-WL-14": Rule(
        rule_id="R-WL-14",
        row_index=15,
        stage="FING&SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.HARD,
        name="特殊首批/质检跟线总成严格限制在指定的受控时间窗口内排产",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-14 独立审计",
        decision_id=None,
        description="特殊首批/质检跟线总成严格限制在指定的受控时间窗口内排产",
        ai_logic="【逻辑公式】特定首批跟线时间窗硬约束：Schedule_Range ⊆ Defined_Time_Slot。\n【业务边界】需特殊资源（质检跟线、设备调机）配套的批次，严格限制在指定时间段。\n【场景示例】指定某高端总成只能在周二下午13:00-17:00（质保主管在岗）生产，排程引擎在该时段独占产线排产。",
    ),
    "R-WL-15": Rule(
        rule_id="R-WL-15",
        row_index=16,
        stage="SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.SOFT,
        name="优先按标准最小包装量（MPQ）整倍数批量排产投料（权重25%）",
        weight=0.25,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-15 独立审计",
        decision_id=None,
        description="优先按标准最小包装量（MPQ）整倍数批量排产投料（权重25%）",
        ai_logic="【逻辑公式】经济批量对齐软约束（权重25%）：Batch_Qty = k × MPQ (k ∈ ℕ⁺)。\n【业务边界】尽量按照标准整批/整箱投料，减少散数、零头料管理成本。\n【场景示例】最小包装量MPQ=500件。需求为650件时，算法权衡凑整为1000件或按650件排产，施加0.25权重约束引导批量规范化。",
    ),
    "R-WL-16": Rule(
        rule_id="R-WL-16",
        row_index=17,
        stage="FG&SFIN",
        category=ConstraintCategory.TEST_SCENARIO,
        nature=ConstraintNature.TEST_CASE,
        name="验证极度空闲下是否正确集中在周一至四排产，且周四结存全覆盖周末需求",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-16 独立审计",
        decision_id=None,
        description="验证极度空闲下是否正确集中在周一至四排产，且周四结存全覆盖周末需求",
        ai_logic="【测试验证】极度空闲用例：输入月需求3000件（极限产能10000件）。\n【验证准则】①排产日全部落于周一至周四；②周五至周日开机工时为0；③周四24:00库存值精确覆盖周五至周日的所有发货量。\n【输出断言】Assert(Days_Worked == Mon_Thu) && Assert(Weekend_Stock_Cover == 100%)。",
    ),
    "R-WL-17": Rule(
        rule_id="R-WL-17",
        row_index=18,
        stage="FG&SFIN",
        category=ConstraintCategory.TEST_SCENARIO,
        nature=ConstraintNature.TEST_CASE,
        name="验证单周需求突发激增时，系统能否将该周动态升阶为单休或连续生产",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-17 独立审计",
        decision_id=None,
        description="验证单周需求突发激增时，系统能否将该周动态升阶为单休或连续生产",
        ai_logic="【测试验证】局部脉冲用例：全月总负荷65%（整体做五休二），但第2周突发激增单周负荷130%。\n【验证准则】①系统第1、3、4周保持做五休二；②仅将第2周局部升阶为做六休一或连续排产；③突发需求消除后自动回退至做五休二。\n【输出断言】Assert(Week2_Calendar == Overtime_7Days) && Assert(Week3_Calendar == Standard_5Days)。",
    ),
    "R-WL-18": Rule(
        rule_id="R-WL-18",
        row_index=19,
        stage="FG&SFIN",
        category=ConstraintCategory.TEST_SCENARIO,
        nature=ConstraintNature.TEST_CASE,
        name="验证前半月爆单时拉满极限产能，绝不盲目均摊，并对溢出订单提前预排或告警",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-18 独立审计",
        decision_id=None,
        description="验证前半月爆单时拉满极限产能，绝不盲目均摊，并对溢出订单提前预排或告警",
        ai_logic="【测试验证】前端拥堵用例：输入前15天负荷150%，后15天负荷20%。\n【验证准则】①禁止在后15天均摊排产导致前半月断货；②前半月拉满7×24H极限产能；③剩余溢出订单向前倒排或精准触发交期延期告警。\n【输出断言】Assert(H1_Utilization == 100%) && Assert(Delay_Notice_Generated == True)。",
    ),
    "R-WL-19": Rule(
        rule_id="R-WL-19",
        row_index=20,
        stage="FG&SFIN",
        category=ConstraintCategory.TEST_SCENARIO,
        nature=ConstraintNature.TEST_CASE,
        name="验证当总需求超出连续生产理论极限产能时，系统直接抛出产能瓶颈缺口告警",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-19 独立审计",
        decision_id=None,
        description="验证当总需求超出连续生产理论极限产能时，系统直接抛出产能瓶颈缺口告警",
        ai_logic="【测试验证】绝对超载用例：全月订单需求 > 全月30天7×24H极限理论产能。\n【验证准则】①算法终止无效寻优；②直接计算绝对产能差额 Gap = Total_Demand - Max_Month_Cap；③精准输出《产能瓶颈告警及委外/扩产建议清单》。\n【输出断言】Assert(Status == \"CAPACITY_OVERLOAD_EXCEPTION\") && Assert(Gap_Qty > 0)。",
    ),
    "R-WL-20": Rule(
        rule_id="R-WL-20",
        row_index=21,
        stage="FING",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.HARD,
        name="当天客户发货需求必须在早班完工入库后全额满足，禁止押后",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-20 独立审计",
        decision_id=None,
        description="当天客户发货需求必须在早班完工入库后全额满足，禁止押后",
        ai_logic="【逻辑公式】早班交付截单硬约束：Morning_Shift_Output + 初始库存 ≥ 当天客户发货总需求。\n【业务边界】成品当日发运以早班生产入库为截止线，早班下班后必须备齐发货量。\n【场景示例】客户周四提货500件。早班初存100件，周四早班装配量必须≥400件并完工入库，杜绝依赖中班或晚班追赶当天发运。",
    ),
    "R-WL-21": Rule(
        rule_id="R-WL-21",
        row_index=22,
        stage="FING&SFIN",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.HARD,
        name="工单最早开工时间受限于BOM关键物料预计入库时间，物料齐套方可开工",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-21 独立审计",
        decision_id=None,
        description="工单最早开工时间受限于BOM关键物料预计入库时间，物料齐套方可开工",
        ai_logic="【逻辑公式】物料齐套开工硬约束：Task_Start(i) ≥ Max { Material_Available_Time(m) | m ∈ BOM(i) }。\n【业务边界】工单开工必须100%齐套，物料未入库/未检验放行前禁止排入生产。\n【场景示例】某总成工单需关键注塑件与芯片。注塑件预计周三10:00入库，芯片预计周三14:00到货，则总成工单最早排产开工时间自动锁定为周三14:00之后。",
    ),
    "R-WL-22": Rule(
        rule_id="R-WL-22",
        row_index=23,
        stage="FING",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.SOFT,
        name="平准化调控排产，控制成品月末安全库存不超过2天日均发货量（权重20%）",
        weight=0.2,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-22 独立审计",
        decision_id=None,
        description="平准化调控排产，控制成品月末安全库存不超过2天日均发货量（权重20%）",
        ai_logic="【逻辑公式】月末安全库存管控软约束（权重20%）：Month_End_Stock ≤ 2 × 日均发货需求。\n【业务边界】在月初及月中产能充裕时，平准化调节日产量，月末库存严禁超过2天。\n【场景示例】日发货量100件，月末安全库存限额200件。系统在月末前夕主动削减高库存物料的日投产，确保月末结存不超200件。",
    ),
    "R-WL-23": Rule(
        rule_id="R-WL-23",
        row_index=24,
        stage="SFIN",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.SOFT,
        name="半成品月末结余库存不得超过设定的安全库存标准天数（权重35%）",
        weight=0.35,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-23 独立审计",
        decision_id=None,
        description="半成品月末结余库存不得超过设定的安全库存标准天数（权重35%）",
        ai_logic="【逻辑公式】月末库存压降软约束（权重35%）：Month_End_Stock ≤ Target_Days × 日均需求。\n【业务边界】月底严格限制半成品沉淀，控制库存占用标准天数。\n【场景示例】安全库存上限为3天用量（300件）。月末排产推演若达420件（超库120件），扣除0.35权重罚分，算法自动减少月末最后两天的投产量。",
    ),
    "R-WL-24": Rule(
        rule_id="R-WL-24",
        row_index=25,
        stage="FING",
        category=ConstraintCategory.ORDER_DELIVERY,
        nature=ConstraintNature.HARD,
        name="向前合批拉单受订单类型严格限制（正式订单≤30天/准度预测≤10天/波动预测≤3-5天）",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-24 独立审计",
        decision_id=None,
        description="向前合批拉单受订单类型严格限制（正式订单≤30天/准度预测≤10天/波动预测≤3-5天）",
        ai_logic="【逻辑公式】需求分级拉单时间窗硬约束：Pull_Window(Order_i) ≤ Max_Advance_Days(Order_Type)；其中正式订单 ≤ 30天，高准度预测 ≤ 10天，波动预测 ≤ 3~5天。\n【业务边界】向前合并批次时严格限制提前期上限，规避低准度预测过早投产导致的死库与缺件风险。\n【场景示例】某产品有下月正式订单和下月波动预测。系统合批时允许拉取30天内的正式订单合并生产，但对波动预测严格限制在5天内，杜绝远期不稳定预测提前生产。",
    ),
    "R-WL-25": Rule(
        rule_id="R-WL-25",
        row_index=26,
        stage="FG&SFIN",
        category=ConstraintCategory.SYSTEM_ARCHITECTURE,
        nature=ConstraintNature.HARD,
        name="产线与工作日历数据彻底解耦，产线按需动态绑定不同日历模板",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-25 独立审计",
        decision_id=None,
        description="产线与工作日历数据彻底解耦，产线按需动态绑定不同日历模板",
        ai_logic="【架构规范】产线-日历解耦：Line_Master 与 Calendar_Template 保持 N:N 弱关联，通过 [Line_ID, Calendar_ID, 生效时段] 动态绑定。\n【业务边界】严禁在算法核心硬编码 weekend_off=True；日历作为外部可配置数据驱动。\n【架构示例】数据库中创建独立日历库（标准双休、做六休一、四班三倒、停线检修）。算法引擎每次调度根据订单负荷动态调用不同日历ID，不修改底层代码。",
    ),
    "R-WL-26": Rule(
        rule_id="R-WL-26",
        row_index=27,
        stage="FG&SFIN",
        category=ConstraintCategory.EQUIPMENT_CAPACITY,
        nature=ConstraintNature.HARD,
        name="按照开班日历排产，班次产能利用率必须保持在95%~105%之间",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-26 独立审计",
        decision_id=None,
        description="按照开班日历排产，班次产能利用率必须保持在95%~105%之间",
        ai_logic="【逻辑公式】班次利用率区间硬约束：95.0% ≤ Shift_Utilization = (实际排产工时 / 班次额定工时) ≤ 105.0%。\n【业务边界】WE34成品装配开班必须维持高效均衡负荷，严禁低于95%低效开机或超过105%过度超载。\n【场景示例】装配线单班额定产能1000件（合规区间950~1050件）。若排产仅800件（利用率80%）或高达1150件（利用率115%），判定违背硬约束无法下发，系统触发合班或平移调整。",
    ),
    "R-WL-27": Rule(
        rule_id="R-WL-27",
        row_index=28,
        stage="FG&SFIN",
        category=ConstraintCategory.EQUIPMENT_CAPACITY,
        nature=ConstraintNature.HARD,
        name="设备计划性预防维护（PM）期间产能为0，排产任务必须避让、切分或顺延",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-WL-27 独立审计",
        decision_id=None,
        description="设备计划性预防维护（PM）期间产能为0，排产任务必须避让、切分或顺延",
        ai_logic="【逻辑公式】维保避让中断硬约束：[Task_Start(k), Task_End(k)] ∩ PM_Window = ∅；若任务跨越维保期，自动切分工单或整体向后顺延。\n【业务边界】计划性保养(PM)窗口设备产能清零，禁止任何排产任务占用。\n【场景示例】机台周四13:00-15:00安排定检保养(PM)。原定12:00-16:00生产的4小时批次，系统自动切分为12:00-13:00(前半段)+15:00-18:00(后半段)或整体延至15:00后开工。",
    ),
}

REGISTRY_SY38_FALLBACK: Dict[str, Rule] = {
    "R-SY-01": Rule(
        rule_id="R-SY-01",
        row_index=29,
        stage="FING",
        category=ConstraintCategory.EXCEPTION_RESPONSE,
        nature=ConstraintNature.HARD,
        name="紧急加急订单特批允许插入当天白班8小时内完工并发货",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-01 独立审计",
        decision_id=None,
        description="紧急加急订单特批允许插入当天白班8小时内完工并发货",
        ai_logic="【逻辑公式】日内紧急响应硬约束：若订单标记为当日加急，允许插入当天白班(08:00-16:00)，完工时间 T_finish ≤ 发货截单时间。\n【业务边界】特批白班8小时快速流转，豁免提前1天备货规则。\n【场景示例】周一早08:30客户突加200件紧急订单需当日17:00发车。系统强制插入周一白班排程并在15:30完工入库，保障准时出货。",
    ),
    "R-SY-02": Rule(
        rule_id="R-SY-02",
        row_index=30,
        stage="SFIN",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.HARD,
        name="保证装配不断线为最高优先级，允许偏撇物料单侧超最大库存生产",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-02 独立审计",
        decision_id=None,
        description="保证装配不断线为最高优先级，允许偏撇物料单侧超最大库存生产",
        ai_logic="【逻辑公式】Penalty(断线) = ∞；若 min(I_L, I_R) < 装配需求，允许 max(I_L, I_R) > 最大库存阈值。\n【业务边界】装配防断线为绝对第一优先，打破库存上限约束。\n【场景示例】装配线需L/R各100件，现有L:20、R:120（最大库存100）。为补足L缺口80件，共模产出80对，导致R库存激增至200，系统判定合法并放行。",
    ),
    "R-SY-03": Rule(
        rule_id="R-SY-03",
        row_index=31,
        stage="SFIN",
        category=ConstraintCategory.TOOLING_PERSONNEL,
        nature=ConstraintNature.HARD,
        name="左右件（L/R）共模联产品必须按模具型腔配比1:1同步排产产出",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-03 独立审计",
        decision_id=None,
        description="左右件（L/R）共模联产品必须按模具型腔配比1:1同步排产产出",
        ai_logic="【逻辑公式】同步产出约束：Output(L, t) = Output(R, t) = 模具冲压次数(t) × 型腔配比(1:1)。\n【业务边界】注塑/冲压模具联产品物理绑定，禁止单件排产。\n【场景示例】模具一模两穴(1L+1R)，排产400模次自动生成L件400件与R件400件，不可单独安排L件生产400件而R件为0。",
    ),
    "R-SY-04": Rule(
        rule_id="R-SY-04",
        row_index=32,
        stage="SFIN",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.HARD,
        name="自制半成品可用库存必须至少提前12小时覆盖后道装配计划需求",
        weight=None,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-04 独立审计",
        decision_id=None,
        description="自制半成品可用库存必须至少提前12小时覆盖后道装配计划需求",
        ai_logic="【逻辑公式】安全前置期硬约束：时刻 t 的可用库存 Stock(t) ≥ ∑ (装配计划需求 [t, t+12H])。\n【业务边界】自制件对后道装配的缓冲期不得少于12小时。\n【场景示例】装配线周二08:00-20:00需消耗500件。注塑机必须在周二08:00前完成生产并检验入库，使08:00动态结存≥500件，否则报错触发前置排产。",
    ),
    "R-SY-05": Rule(
        rule_id="R-SY-05",
        row_index=33,
        stage="SFIN",
        category=ConstraintCategory.CORE_SCHEDULING,
        nature=ConstraintNature.SOFT,
        name="单机台12小时内常规最多排产2套模具，极度紧张时允许增产但扣罚",
        weight=0.3,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-05 独立审计",
        decision_id=None,
        description="单机台12小时内常规最多排产2套模具，极度紧张时允许增产但扣罚",
        ai_logic="【逻辑公式】切换频次软约束：Count(换模次数 in 12H) ≤ 1；若换模次数 > 1，施加权重0.30惩罚。\n【业务边界】常规单班最多跑2副模具；产能极度紧缺时允许通过损失工时换模来保交付。\n【场景示例】机台12H班次内已有模具A和B。若插入急单模具C，产生第2次换模（损耗45分钟），算法扣除30%权重惩罚，权衡产出损失与断线风险后决策。",
    ),
    "R-SY-06": Rule(
        rule_id="R-SY-06",
        row_index=34,
        stage="FING",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.SOFT,
        name="成品动态结余库存上限不超过未来5天发货需求总量（权重60%）",
        weight=0.6,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-06 独立审计",
        decision_id=None,
        description="成品动态结余库存上限不超过未来5天发货需求总量（权重60%）",
        ai_logic="【逻辑公式】库龄与库容软约束（权重60%）：Inventory(t) ≤ ∑_{k=1..5} Daily_Demand(t+k)。\n【业务边界】成品最大库存上限为动态向前覆盖5天需求，抑制过量积压。\n【场景示例】未来第1至第5天客户总需求为2500件。若排产使库存达到3100件（超额600件），触发0.60权重超库惩罚，驱动排产推后或降速。",
    ),
    "R-SY-07": Rule(
        rule_id="R-SY-07",
        row_index=35,
        stage="SFIN",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.SOFT,
        name="在保证不断线前提下半成品库存不超上限，超库施加80%惩罚",
        weight=0.8,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-07 独立审计",
        decision_id=None,
        description="在保证不断线前提下半成品库存不超上限，超库施加80%惩罚",
        ai_logic="【逻辑公式】软约束惩罚项：Min ∑ [ 0.80 × OverStock_Cost × Max(0, Stock(t) - Max_Limit) ]。\n【业务边界】库存上限控制，在不断线前提下超库需扣减优化目标得分。\n【场景示例】SFIN件设上限1000件。方案A库存1050件（罚分50×0.8=40），方案B库存950件（罚分0），在均不断线下求解器优先选择方案B。",
    ),
    "R-SY-08": Rule(
        rule_id="R-SY-08",
        row_index=36,
        stage="FING",
        category=ConstraintCategory.EQUIPMENT_CAPACITY,
        nature=ConstraintNature.SOFT,
        name="班次产能利用率目标控制在95%~105%之间，偏离施加90%惩罚",
        weight=0.9,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-08 独立审计",
        decision_id=None,
        description="班次产能利用率目标控制在95%~105%之间，偏离施加90%惩罚",
        ai_logic="【逻辑公式】双边利用率目标（权重90%）：95% ≤ (实际排产工时 / 班次额定工时) ≤ 105%。\n【业务边界】避免产线产能浪费或设备过载过劳，保持生产节奏稳定。\n【场景示例】班次标称产能1000件，合规区间[950, 1050]。若排产800件(利用率80%)或1150件(利用率115%)，优化器按0.90权重施加严重偏离惩罚，推动负荷调整。",
    ),
    "R-SY-09": Rule(
        rule_id="R-SY-09",
        row_index=37,
        stage="FING",
        category=ConstraintCategory.MATERIAL_INVENTORY,
        nature=ConstraintNature.SOFT,
        name="装配完工结余库存须至少提前1天满足次日发货需求（权重90%）",
        weight=0.9,
        status=RuleStatus.IMPLEMENTED,
        landing_layer="排产引擎核心约束层",
        validation_check="规则 R-SY-09 独立审计",
        decision_id=None,
        description="装配完工结余库存须至少提前1天满足次日发货需求（权重90%）",
        ai_logic="【逻辑公式】提前出货储备（权重90%）：Day(N)完工结余库存 ≥ Day(N+1)客户发货计划总量。\n【业务边界】建立1天成品安全缓冲期，防范当日生产扰动导致交付违约。\n【场景示例】周三客户要货800件。装配线在周二24:00前的生产结余必须≥800件，确保周三早间直接装车发运。",
    ),
}


def get_rules_for_site(site: str = "WL", excel_path: Optional[str] = None) -> Dict[str, Rule]:
    """Retrieve rules dynamically from Excel or fallback for the specified factory site."""
    try:
        rules_by_site = load_rules_from_excel(excel_path)
        if site in rules_by_site and rules_by_site[site]:
            return rules_by_site[site]
    except Exception:
        pass
    if site == "SY38":
        return REGISTRY_SY38_FALLBACK
    return REGISTRY_WE34_FALLBACK

# Default backward compatibility registry
REGISTRY: Dict[str, Rule] = REGISTRY_WE34_FALLBACK
